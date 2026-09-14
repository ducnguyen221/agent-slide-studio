from __future__ import annotations

import hashlib
import importlib.util
import json
import os
from pathlib import Path
import _thread
import subprocess
import sys
from threading import Barrier, Thread
import time
from types import SimpleNamespace

import pytest
import yaml
from PIL import Image
from pptx import Presentation

from presentation_studio import cli
from presentation_studio.backends.pptx_native import (
    NativeBuildError,
    PptxNativeBackend,
    RendererDiscovery,
    discover_renderer,
    extract_semantics,
    native_capabilities,
    render_staging_path,
    wait_for_exported_slides,
)
from presentation_studio.models import BuildInput, CLIResult, DeckSpec, ProjectConfig
from presentation_studio.state import hash_inputs
from presentation_studio.workspace import init_project, resolve_paths


FIXTURE = Path(__file__).parents[1] / "fixtures" / "native" / "deck.yaml"


def _blocking_renderer_worker(connection: object, delay_seconds: float) -> None:
    del connection
    time.sleep(delay_seconds)


def _blocking_renderer_worker_with_pid(
    connection: object, worker_pid_file: str, delay_seconds: float
) -> None:
    del connection
    Path(worker_pid_file).write_text(str(os.getpid()), encoding="ascii")
    time.sleep(delay_seconds)


def _baseline_only_renderer_worker(
    connection: object,
    baseline: tuple[int, ...],
    worker_pid_file: str,
    delay_seconds: float,
) -> None:
    Path(worker_pid_file).write_text(str(os.getpid()), encoding="ascii")
    connection.send({"kind": "powerpoint-baseline", "pids": list(baseline)})
    time.sleep(delay_seconds)


def _crashing_renderer_worker(connection: object, worker_pid_file: str) -> None:
    del connection
    Path(worker_pid_file).write_text(str(os.getpid()), encoding="ascii")
    os._exit(23)


def _error_renderer_worker(connection: object, worker_pid_file: str) -> None:
    Path(worker_pid_file).write_text(str(os.getpid()), encoding="ascii")
    connection.send(
        {
            "kind": "error",
            "primary_type": "SyntheticRendererError",
            "primary_message": "synthetic worker failure",
            "notes": ["cleanup evidence"],
        }
    )
    connection.close()


def _pid_is_running(pid: int) -> bool:
    if sys.platform == "win32":
        import ctypes

        kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel.OpenProcess.argtypes = [ctypes.c_uint32, ctypes.c_int, ctypes.c_uint32]
        kernel.OpenProcess.restype = ctypes.c_void_p
        kernel.GetExitCodeProcess.argtypes = [
            ctypes.c_void_p,
            ctypes.POINTER(ctypes.c_uint32),
        ]
        kernel.GetExitCodeProcess.restype = ctypes.c_int
        kernel.CloseHandle.argtypes = [ctypes.c_void_p]
        kernel.CloseHandle.restype = ctypes.c_int
        handle = kernel.OpenProcess(0x1000 | 0x100000, False, pid)
        if not handle:
            error = ctypes.get_last_error()
            if error == 87:
                return False
            raise ctypes.WinError(error)
        try:
            exit_code = ctypes.c_uint32()
            if not kernel.GetExitCodeProcess(handle, ctypes.byref(exit_code)):
                raise ctypes.WinError(ctypes.get_last_error())
            return exit_code.value == 259
        finally:
            kernel.CloseHandle(handle)
    try:
        os.kill(pid, 0)
    except OSError:
        return False
    return True


class _TerminateRaisesWorker:
    def __init__(self) -> None:
        self.alive = True
        self.calls: list[str] = []

    def is_alive(self) -> bool:
        return self.alive

    def terminate(self) -> None:
        self.calls.append("terminate")
        raise OSError("synthetic terminate failure")

    def kill(self) -> None:
        self.calls.append("kill")
        self.alive = False

    def join(self, timeout: float) -> None:
        self.calls.append("join")


class _TerminateRaisesExternal:
    def __init__(self) -> None:
        self.alive = True
        self.calls: list[str] = []

    def poll(self) -> int | None:
        return None if self.alive else -9

    def terminate(self) -> None:
        self.calls.append("terminate")
        raise OSError("synthetic external terminate failure")

    def kill(self) -> None:
        self.calls.append("kill")
        self.alive = False

    def wait(self, timeout: float) -> int:
        self.calls.append("wait")
        if self.alive:
            raise subprocess.TimeoutExpired("synthetic", timeout)
        return -9


def _run_supervisor_with_post_start_failure(
    monkeypatch: pytest.MonkeyPatch, primary_error: BaseException
) -> tuple[BaseException, bool]:
    from presentation_studio.backends import pptx_native as native

    process_type = type(native.multiprocessing.get_context("spawn").Process())
    original_start = process_type.start
    started_processes: list[object] = []

    def start_then_raise(process: object) -> None:
        original_start(process)
        started_processes.append(process)
        raise primary_error

    monkeypatch.setattr(process_type, "start", start_then_raise)
    caught: BaseException | None = None
    worker_was_alive = True
    try:
        try:
            native._run_owned_renderer_process(
                _blocking_renderer_worker,
                (30.0,),
                timeout_seconds=2,
            )
        except BaseException as error:
            caught = error
        assert len(started_processes) == 1
        worker_pid = started_processes[0].pid
        assert worker_pid is not None
        worker_was_alive = _pid_is_running(worker_pid)
    finally:
        for process in started_processes:
            if process.is_alive():
                process.kill()
            process.join(2)

    assert caught is not None
    return caught, worker_was_alive


def _deck() -> DeckSpec:
    return DeckSpec.model_validate(yaml.safe_load(FIXTURE.read_text(encoding="utf-8")))


def _build_input(deck: DeckSpec) -> BuildInput:
    project = ProjectConfig(project_id="native-demo", title="Báo cáo vận hành minh họa")
    payload = {
        "project": project.model_dump(mode="json"),
        "deck": deck.model_dump(mode="json"),
    }
    return BuildInput(project=project, deck=deck, input_hash=hash_inputs(payload))


def _backend_without_renderer() -> PptxNativeBackend:
    return PptxNativeBackend(
        renderer=RendererDiscovery(
            renderer_id="none",
            available=False,
            unavailable_reason="Không tìm thấy renderer trong môi trường kiểm thử.",
        )
    )


def test_build_fails_before_staging_when_identity_safe_promotion_is_unavailable(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from presentation_studio.backends import pptx_native as native

    paths = resolve_paths(workspace=tmp_path)
    monkeypatch.setattr(
        native,
        "_identity_safe_native_promotion_available",
        lambda: False,
        raising=False,
    )

    with pytest.raises(RuntimeError, match="identity-safe"):
        _backend_without_renderer().build(_build_input(_deck()), paths)

    assert not paths.builds_dir.exists()


def test_capability_matrix_fails_closed_without_identity_safe_promotion(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from presentation_studio.backends import pptx_native as native

    monkeypatch.setattr(
        native, "_identity_safe_native_promotion_available", lambda: False
    )

    capabilities = native.native_capabilities(
        RendererDiscovery(renderer_id="none", available=False)
    )

    assert capabilities.available is False
    assert capabilities.unavailable_reason is not None
    assert "identity-safe" in capabilities.unavailable_reason


def test_capability_matrix_is_explicit_and_renderer_status_is_honest() -> None:
    discovery = RendererDiscovery(
        renderer_id="none",
        available=False,
        unavailable_reason="Không có renderer.",
    )

    capabilities = native_capabilities(discovery)

    assert capabilities.available is True
    assert capabilities.supported_inputs == ["semantic"]
    assert capabilities.supported_outputs == ["pptx"]
    assert capabilities.supported_elements == [
        "chart",
        "process",
        "quote",
        "table",
        "text",
        "timeline",
    ]
    assert capabilities.editable_elements == capabilities.supported_elements
    assert capabilities.notes_support_by_output == {"pptx": True}
    assert capabilities.requires_network is False
    assert capabilities.render_capabilities.png is False
    assert capabilities.verification_by_feature["visual-render"].status == "unverified"
    assert capabilities.verification_by_feature["visual-render"].evidence_ref is None


def test_renderer_discovery_reports_missing_pywin32_as_unavailable(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def missing_parent(name: str) -> object:
        raise ModuleNotFoundError(name)

    monkeypatch.setattr(importlib.util, "find_spec", missing_parent)

    discovery = discover_renderer()

    assert discovery.available is False
    assert discovery.renderer_id == "none"
    assert discovery.unavailable_reason == "Thiếu pywin32 để điều khiển PowerPoint renderer."


def test_render_waits_for_async_powerpoint_export(tmp_path: Path) -> None:
    def delayed_export() -> None:
        time.sleep(0.03)
        (tmp_path / "Slide1.PNG").write_bytes(b"first")
        (tmp_path / "Slide2.PNG").write_bytes(b"second")

    writer = Thread(target=delayed_export)
    writer.start()
    try:
        exported = wait_for_exported_slides(
            tmp_path, expected_count=2, timeout_seconds=0.5
        )
    finally:
        writer.join()

    assert [path.name for path in exported] == ["Slide1.PNG", "Slide2.PNG"]


def test_render_waits_until_powerpoint_releases_temporary_hardlink(
    tmp_path: Path,
) -> None:
    rendered = tmp_path / "Slide1.PNG"
    temporary_link = tmp_path / "powerpoint-export.tmp"
    rendered.write_bytes(b"stable-size")
    os.link(rendered, temporary_link)

    def release_export_link() -> None:
        time.sleep(0.2)
        temporary_link.unlink()

    writer = Thread(target=release_export_link)
    writer.start()
    started = time.monotonic()
    returned_at = started
    try:
        exported = wait_for_exported_slides(
            tmp_path, expected_count=1, timeout_seconds=1
        )
        returned_at = time.monotonic()
    finally:
        writer.join()

    assert returned_at - started >= 0.15
    assert exported[0].stat().st_nlink == 1


def test_powerpoint_render_staging_path_has_no_suffix_for_export_reliability(
    tmp_path: Path,
) -> None:
    # PowerPoint Export strips a suffix such as .tmp and writes to a sibling path.
    staging = render_staging_path(tmp_path, "abc123")

    assert staging == tmp_path / "staging-abc123"
    assert staging.suffix == ""


def test_native_build_reopens_and_preserves_semantics_order_notes_and_data(
    tmp_path: Path,
) -> None:
    paths = resolve_paths(workspace=tmp_path)
    deck = _deck()

    result = _backend_without_renderer().build(_build_input(deck), paths)

    assert result.status == "passed"
    assert result.expected_slide_ids == ["tong-quan", "bang-du-lieu", "bieu-do"]
    assert result.actual_slide_ids == result.expected_slide_ids
    assert [item.slide_id for item in result.slide_results] == result.expected_slide_ids
    assert result.slide_results[0].element_editability == {
        "e-kpi": "native",
        "e-explain": "native",
    }
    assert result.slide_results[1].element_editability == {"e-table": "native"}
    assert result.slide_results[2].element_editability == {"e-chart": "native"}
    assert result.warnings == ["Chưa kiểm chứng hình ảnh: Không tìm thấy renderer trong môi trường kiểm thử."]

    output = tmp_path / result.outputs[0].path
    assert output.is_file()
    assert hashlib.sha256(output.read_bytes()).hexdigest() == result.outputs[0].sha256

    semantic = extract_semantics(output)
    assert [slide.slide_id for slide in semantic] == result.expected_slide_ids
    assert [slide.notes for slide in semantic] == [slide.notes for slide in deck.slides]
    assert semantic[0].source_refs == ["demo-data"]
    assert semantic[0].elements["e-kpi"].text == "18%"
    assert semantic[0].elements["e-explain"].text == "Hai thay đổi quy trình giúp giảm thời gian chờ."
    assert semantic[1].elements["e-table"].columns == ["Nhóm", "Trước", "Sau"]
    assert semantic[1].elements["e-table"].rows == [
        ["A", "25", "18"],
        ["B", "30", "21"],
    ]
    assert semantic[2].elements["e-chart"].chart_type == "column"
    assert semantic[2].elements["e-chart"].data == [("Tháng 1", 72.0), ("Tháng 2", 81.0)]
    assert semantic[2].elements["e-chart"].unit == "điểm"
    assert semantic[2].elements["e-chart"].source_ref == "demo-data"
    presentation = Presentation(output)
    explanation = next(
        shape
        for shape in presentation.slides[0].shapes
        if shape.name == "presentation-studio:element:e-explain:text"
    )
    assert explanation._element.xpath(".//p:cNvPr")[0].get("descr") == (
        "Giải thích hai thay đổi quy trình"
    )


def test_native_textbox_remains_editable_after_reopen(tmp_path: Path) -> None:
    paths = resolve_paths(workspace=tmp_path)
    result = _backend_without_renderer().build(_build_input(_deck()), paths)
    output = tmp_path / result.outputs[0].path
    presentation = Presentation(output)
    shape = next(
        shape
        for shape in presentation.slides[0].shapes
        if shape.name == "presentation-studio:element:e-explain:text"
    )
    shape.text = "Nội dung đã sửa bằng python-pptx."
    presentation.save(output)

    semantic = extract_semantics(output)

    assert semantic[0].elements["e-explain"].text == "Nội dung đã sửa bằng python-pptx."


def test_custom_canvas_keeps_every_native_shape_inside_slide(tmp_path: Path) -> None:
    deck = DeckSpec(
        title="Canvas tùy chỉnh",
        audience="A",
        purpose="P",
        canvas={"ratio": "custom", "width": 10, "height": 5.625, "unit": "inch"},
        slides=[
            {
                "slide_id": "s01",
                "title": "Nội dung nằm trong canvas",
                "message": "Không có đối tượng bị đặt ngoài slide.",
                "layout_ref": "builtin:content@1.0.0",
                "elements": [
                    {"element_id": "e01", "kind": "text", "content": "Nội dung"}
                ],
            }
        ],
    )
    paths = resolve_paths(workspace=tmp_path)

    result = _backend_without_renderer().build(_build_input(deck), paths)
    presentation = Presentation(tmp_path / result.outputs[0].path)
    slide = presentation.slides[0]

    assert presentation.slide_width == pytest.approx(10 * 914400, abs=2)
    assert presentation.slide_height == pytest.approx(5.625 * 914400, abs=2)
    assert all(shape.left >= 0 and shape.top >= 0 for shape in slide.shapes)
    assert all(shape.left + shape.width <= presentation.slide_width for shape in slide.shapes)
    assert all(shape.top + shape.height <= presentation.slide_height for shape in slide.shapes)


@pytest.mark.parametrize(
    "deck",
    [
        DeckSpec(
            title="Sai layout",
            audience="A",
            purpose="P",
            slides=[
                {
                    "slide_id": "s01",
                    "title": "Không được fallback",
                    "layout_ref": "builtin:unknown@1.0.0",
                }
            ],
        ),
        DeckSpec(
            title="Quá sức chứa",
            audience="A",
            purpose="P",
            slides=[
                {
                    "slide_id": "s01",
                    "title": "Không được cắt",
                    "layout_ref": "builtin:timeline@1.0.0",
                    "elements": [
                        {
                            "element_id": "e-timeline",
                            "kind": "timeline",
                            "content": {
                                "items": [
                                    {"time_label": str(i), "title": f"Mốc {i}"}
                                    for i in range(5)
                                ]
                            },
                        }
                    ],
                }
            ],
        ),
    ],
)
def test_preflight_failure_does_not_promote_partial_output(
    tmp_path: Path, deck: DeckSpec
) -> None:
    paths = resolve_paths(workspace=tmp_path)

    with pytest.raises(NativeBuildError) as raised:
        _backend_without_renderer().build(_build_input(deck), paths)

    assert raised.value.issues
    assert not list(tmp_path.rglob("*.pptx"))
    assert not [path for path in paths.builds_dir.glob("*") if not path.name.startswith(".")]


def test_cli_build_returns_one_honest_json_object(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    paths = resolve_paths(workspace=tmp_path)
    init_project(paths, project_id="native-demo", title="Báo cáo vận hành minh họa")
    (tmp_path / "storyboard" / "deck.yaml").write_text(
        FIXTURE.read_text(encoding="utf-8"), encoding="utf-8"
    )
    cli._BACKENDS.clear()
    cli._RENDERERS.clear()

    exit_code = cli.run(
        [
            "build",
            "--workspace",
            str(tmp_path),
            "--backend",
            "pptx-native",
            "--json",
        ]
    )
    captured = capsys.readouterr()
    payload = json.loads(captured.out)

    assert exit_code == 0
    assert len([line for line in captured.out.splitlines() if line.strip()]) == 1
    assert captured.err == ""
    assert payload["command"] == "build"
    assert payload["status"] in {"passed", "unverified"}
    assert payload["data"]["build"]["status"] == "passed"
    assert payload["data"]["visual_status"] in {"passed", "unverified"}
    assert payload["data"]["capabilities"]["verification_by_feature"][
        "semantic-reopen"
    ]["status"] == "passed"
    assert payload["data"]["capabilities"]["verification_by_feature"][
        "semantic-reopen"
    ]["evidence_ref"] == payload["data"]["build"]["outputs"][0]["path"]
    assert (tmp_path / payload["data"]["build"]["outputs"][0]["path"]).is_file()


def test_builtin_registration_does_not_replace_explicit_backend(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    cli._BACKENDS.clear()
    cli._RENDERERS.clear()
    cli.register_backend(
        "pptx-native",
        lambda paths: CLIResult(
            command="build",
            status="passed",
            exit_code=0,
            data={"handler": "explicit"},
        ),
    )

    exit_code = cli.run(
        [
            "build",
            "--workspace",
            str(tmp_path),
            "--backend",
            "pptx-native",
            "--json",
        ]
    )
    payload = json.loads(capsys.readouterr().out)

    assert exit_code == 0
    assert payload["data"] == {"handler": "explicit"}


def test_render_keeps_raw_export_names_separate_from_slide_ids(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from presentation_studio.backends import pptx_native as native

    paths = resolve_paths(workspace=tmp_path)
    build_id = "collision-build"
    source = tmp_path / "builds" / build_id / "deck.pptx"
    source.parent.mkdir(parents=True)
    source.write_bytes(b"fixture")
    monkeypatch.setattr(
        native,
        "extract_semantics",
        lambda path: [
            SimpleNamespace(slide_id="slide2"),
            SimpleNamespace(slide_id="slide1"),
        ],
    )

    def export_runner(source_path, raw_directory, expected_count, timeout_seconds):
        del source_path, timeout_seconds
        assert expected_count == 2
        Image.new("RGB", (40, 30), "red").save(raw_directory / "Slide1.PNG")
        Image.new("RGB", (40, 30), "blue").save(raw_directory / "Slide2.PNG")

    backend = PptxNativeBackend(
        renderer=RendererDiscovery(renderer_id="powerpoint-com", available=True),
        renderer_runner=export_runner,
    )

    result = backend.render(paths, build_id, timeout_seconds=1)

    assert list(result.slides) == ["slide2", "slide1"]
    actual_hashes = []
    for rendered in result.slides.values():
        artifact = tmp_path / rendered.path
        assert artifact.is_file()
        digest = hashlib.sha256(artifact.read_bytes()).hexdigest()
        assert digest == rendered.sha256
        actual_hashes.append(digest)
    assert len(set(actual_hashes)) == 2


def test_bound_staging_cleanup_refuses_foreign_path_replacement(tmp_path: Path) -> None:
    from presentation_studio.backends.pptx_native import _BoundStage

    parent = tmp_path / "parent"
    parent.mkdir()
    stage = _BoundStage.create(parent, "stage-owned")
    saved_owned = parent / "saved-owned"
    try:
        stage.directory.close()
        stage.path.rename(saved_owned)
        stage.path.mkdir()
        marker = stage.path / "foreign.txt"
        marker.write_text("foreign", encoding="utf-8")

        with pytest.raises(OSError):
            stage.remove()

        assert marker.read_text(encoding="utf-8") == "foreign"
        assert saved_owned.is_dir()
    finally:
        stage.close()


def test_staging_creation_never_adopts_foreign_directory_from_create_race(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from presentation_studio.backends import pptx_native as native
    from presentation_studio.fs import BoundDirectory

    parent = tmp_path / "builds"
    parent.mkdir()
    foreign = parent / "staging-race"
    foreign_marker = foreign / "foreign.txt"
    original_child = BoundDirectory.child

    def create_foreign_then_delegate(self, name, **kwargs):
        if kwargs.get("create") and name == "staging-race" and not foreign.exists():
            foreign.mkdir()
            foreign_marker.write_text("foreign", encoding="utf-8")
        return original_child(self, name, **kwargs)

    monkeypatch.setattr(BoundDirectory, "child", create_foreign_then_delegate)
    stage = None
    error = None
    try:
        stage = native._BoundStage.create(parent, "staging-race")
        stage.remove()
    except BaseException as exc:
        error = exc
    finally:
        if stage is not None:
            try:
                stage.close()
            except OSError:
                pass

    assert error is not None, "exclusive staging creation accepted a foreign directory"
    assert foreign_marker.read_text(encoding="utf-8") == "foreign"


@pytest.mark.skipif(sys.platform != "win32", reason="Windows atomic directory handle")
@pytest.mark.parametrize("fault", ["lstat", "identity", "expected"])
def test_exclusive_child_closes_atomic_descriptor_after_validation_failure(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, fault: str
) -> None:
    from presentation_studio import fs

    parent_path = tmp_path / "parent"
    parent_path.mkdir()
    decoy_file = tmp_path / "decoy.txt"
    decoy_file.write_text("not a directory", encoding="utf-8")
    expected_dir = tmp_path / "expected"
    expected_dir.mkdir()
    parent = fs.BoundDirectory.open(parent_path)
    original_create = fs._windows_create_directory_at
    original_lstat = parent.lstat
    descriptors: list[int] = []

    def capture_descriptor(*args, **kwargs):
        descriptor = original_create(*args, **kwargs)
        descriptors.append(descriptor)
        return descriptor

    def inject_lstat(name: str):
        if name == "candidate" and descriptors:
            if fault == "lstat":
                raise FileNotFoundError("synthetic post-create lstat failure")
            if fault == "identity":
                return decoy_file.lstat()
        return original_lstat(name)

    monkeypatch.setattr(fs, "_windows_create_directory_at", capture_descriptor)
    monkeypatch.setattr(parent, "lstat", inject_lstat)
    expected = expected_dir.lstat() if fault == "expected" else None
    try:
        with pytest.raises((FileNotFoundError, fs.UnsafeFileError)):
            parent.child(
                "candidate",
                create=True,
                exclusive=True,
                allow_delete=True,
                expected=expected,
            )
        assert descriptors
        with pytest.raises(OSError):
            os.fstat(descriptors[0])
    finally:
        for descriptor in descriptors:
            try:
                os.close(descriptor)
            except OSError:
                pass
        parent.close()


@pytest.mark.skipif(sys.platform != "win32", reason="Windows atomic directory handle")
def test_exclusive_child_preserves_primary_when_descriptor_close_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from presentation_studio import fs

    parent_path = tmp_path / "parent"
    parent_path.mkdir()
    parent = fs.BoundDirectory.open(parent_path)
    original_create = fs._windows_create_directory_at
    original_lstat = parent.lstat
    real_close = os.close
    descriptors: list[int] = []
    close_attempts = 0

    def capture_descriptor(*args, **kwargs):
        descriptor = original_create(*args, **kwargs)
        descriptors.append(descriptor)
        return descriptor

    def fail_lstat(name: str):
        if name == "candidate" and descriptors:
            raise RuntimeError("primary identity validation failure")
        return original_lstat(name)

    def fail_descriptor_close(descriptor: int) -> None:
        nonlocal close_attempts
        if descriptors and descriptor == descriptors[0] and close_attempts == 0:
            close_attempts += 1
            raise OSError("secondary descriptor close failure")
        real_close(descriptor)

    monkeypatch.setattr(fs, "_windows_create_directory_at", capture_descriptor)
    monkeypatch.setattr(parent, "lstat", fail_lstat)
    monkeypatch.setattr(os, "close", fail_descriptor_close)
    try:
        with pytest.raises(RuntimeError, match="primary identity") as raised:
            parent.child(
                "candidate", create=True, exclusive=True, allow_delete=True
            )
        assert close_attempts == 1
        assert any("close" in note and "OSError" in note for note in raised.value.__notes__)
    finally:
        monkeypatch.setattr(os, "close", real_close)
        for descriptor in descriptors:
            try:
                real_close(descriptor)
            except OSError:
                pass
        parent.close()


def test_build_safe_write_never_overwrites_foreign_hardlink(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from pptx.presentation import Presentation as PresentationClass

    paths = resolve_paths(workspace=tmp_path)
    foreign = tmp_path / "foreign.bin"
    original_bytes = b"foreign-bytes-must-survive"
    foreign.write_bytes(original_bytes)
    original_save = PresentationClass.save
    linked = False

    def install_hardlink_then_save(self, target) -> None:
        nonlocal linked
        staging = next(paths.builds_dir.glob("staging-*"))
        candidate = staging / "deck.pptx"
        os.link(foreign, candidate)
        linked = True
        original_save(self, target)

    monkeypatch.setattr(PresentationClass, "save", install_hardlink_then_save)

    with pytest.raises(BaseException):
        _backend_without_renderer().build(_build_input(_deck()), paths)

    assert linked is True
    assert foreign.read_bytes() == original_bytes
    assert not [path for path in paths.builds_dir.glob("native-*")]


@pytest.mark.parametrize(
    ("fault_name", "fault_type"),
    [("raw", OSError), ("mapped", ValueError), ("mapped", KeyboardInterrupt)],
)
def test_render_acquisition_failure_cleans_every_resource_acquired_so_far(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    fault_name: str,
    fault_type: type[BaseException],
) -> None:
    from presentation_studio.backends import pptx_native as native
    from presentation_studio.fs import BoundDirectory

    paths = resolve_paths(workspace=tmp_path)
    build_id = "acquisition-build"
    source = tmp_path / "builds" / build_id / "deck.pptx"
    source.parent.mkdir(parents=True)
    source.write_bytes(b"fixture")
    monkeypatch.setattr(
        native,
        "extract_semantics",
        lambda path: [SimpleNamespace(slide_id="s01")],
    )
    original_child = BoundDirectory.child
    original_close = BoundDirectory.close
    original_create = native._BoundStage.create
    original_remove = native._BoundStage.remove
    original_stage_close = native._BoundStage.close
    captured_stage = []
    captured_raw = []
    calls = {"raw.close": 0, "stage.remove": 0, "stage.close": 0}

    def capture_create(*args, **kwargs):
        stage = original_create(*args, **kwargs)
        captured_stage.append(stage)
        return stage

    def inject_fault(self, name, **kwargs):
        if name == fault_name:
            raise fault_type(f"synthetic {fault_name} acquisition failure")
        child = original_child(self, name, **kwargs)
        if name == "raw":
            captured_raw.append(child)
        return child

    def track_close(self):
        if captured_raw and self is captured_raw[0]:
            calls["raw.close"] += 1
        return original_close(self)

    def track_remove(self):
        calls["stage.remove"] += 1
        return original_remove(self)

    def track_stage_close(self):
        calls["stage.close"] += 1
        return original_stage_close(self)

    monkeypatch.setattr(native._BoundStage, "create", capture_create)
    monkeypatch.setattr(BoundDirectory, "child", inject_fault)
    monkeypatch.setattr(BoundDirectory, "close", track_close)
    monkeypatch.setattr(native._BoundStage, "remove", track_remove)
    monkeypatch.setattr(native._BoundStage, "close", track_stage_close)
    observed = None
    try:
        with pytest.raises(fault_type, match=f"synthetic {fault_name}"):
            PptxNativeBackend(
                renderer=RendererDiscovery(renderer_id="powerpoint-com", available=True),
                renderer_runner=lambda *args: None,
            ).render(paths, build_id, timeout_seconds=1)
        observed = dict(calls)
    finally:
        for raw in captured_raw:
            original_close(raw)
        for stage in captured_stage:
            if stage.path.exists() and not stage.removed:
                original_remove(stage)
            original_stage_close(stage)

    assert observed is not None
    assert observed["stage.remove"] == 1
    assert observed["stage.close"] == 1
    assert observed["raw.close"] == (0 if fault_name == "raw" else 1)
    assert not [path for path in source.parent.joinpath("renders").glob("staging-*")]


def test_render_acquisition_preserves_primary_when_raw_close_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from presentation_studio.backends import pptx_native as native
    from presentation_studio.fs import BoundDirectory

    paths = resolve_paths(workspace=tmp_path)
    build_id = "acquisition-close-build"
    source = tmp_path / "builds" / build_id / "deck.pptx"
    source.parent.mkdir(parents=True)
    source.write_bytes(b"fixture")
    monkeypatch.setattr(
        native,
        "extract_semantics",
        lambda path: [SimpleNamespace(slide_id="s01")],
    )
    original_child = BoundDirectory.child
    original_close = BoundDirectory.close
    captured_raw = []

    def fail_mapped(self, name, **kwargs):
        if name == "mapped":
            raise ValueError("primary mapped acquisition failure")
        child = original_child(self, name, **kwargs)
        if name == "raw":
            captured_raw.append(child)
        return child

    def close_raw_then_raise(self):
        original_close(self)
        if captured_raw and self is captured_raw[0]:
            raise OSError("secondary raw close failure")

    monkeypatch.setattr(BoundDirectory, "child", fail_mapped)
    monkeypatch.setattr(BoundDirectory, "close", close_raw_then_raise)

    with pytest.raises(ValueError, match="primary mapped") as raised:
        PptxNativeBackend(
            renderer=RendererDiscovery(renderer_id="powerpoint-com", available=True),
            renderer_runner=lambda *args: None,
        ).render(paths, build_id, timeout_seconds=1)

    assert any("raw" in note and "OSError" in note for note in raised.value.__notes__)
    assert not [path for path in source.parent.joinpath("renders").glob("staging-*")]


def test_powerpoint_cleanup_preserves_export_error_and_attempts_all_steps(
    tmp_path: Path,
) -> None:
    from presentation_studio.backends.pptx_native import _powerpoint_export_lifecycle

    calls: list[str] = []

    class FakePresentation:
        def Export(self, destination: str, format_name: str) -> None:
            del destination, format_name
            calls.append("Export")
            raise ValueError("primary export failure")

        def Close(self) -> None:
            calls.append("Close")
            raise OSError("secondary close failure")

    class FakeApplication:
        Presentations = SimpleNamespace(Open=lambda *args, **kwargs: FakePresentation())

        def Quit(self) -> None:
            calls.append("Quit")

    pythoncom = SimpleNamespace(
        CoInitialize=lambda: calls.append("CoInitialize"),
        CoUninitialize=lambda: calls.append("CoUninitialize"),
    )

    with pytest.raises(ValueError, match="primary export failure") as raised:
        _powerpoint_export_lifecycle(
            tmp_path / "deck.pptx",
            tmp_path / "raw",
            expected_count=1,
            timeout_seconds=1,
            pythoncom_module=pythoncom,
            dispatch_ex=lambda progid: FakeApplication(),
        )

    assert calls == ["CoInitialize", "Export", "Close", "Quit", "CoUninitialize"]
    assert raised.value.__notes__ == ["PowerPoint cleanup failed: Close: OSError"]


def test_powerpoint_cleanup_failure_after_export_is_reported_after_all_steps(
    tmp_path: Path,
) -> None:
    from presentation_studio.backends.pptx_native import (
        RendererCleanupError,
        _powerpoint_export_lifecycle,
    )

    raw = tmp_path / "raw"
    raw.mkdir()
    calls: list[str] = []

    class FakePresentation:
        def Export(self, destination: str, format_name: str) -> None:
            calls.append("Export")
            assert format_name == "PNG"
            Image.new("RGB", (40, 30), "green").save(Path(destination) / "Slide1.PNG")

        def Close(self) -> None:
            calls.append("Close")
            raise OSError("secondary close failure")

    class FakeApplication:
        Presentations = SimpleNamespace(Open=lambda *args, **kwargs: FakePresentation())

        def Quit(self) -> None:
            calls.append("Quit")

    pythoncom = SimpleNamespace(
        CoInitialize=lambda: calls.append("CoInitialize"),
        CoUninitialize=lambda: calls.append("CoUninitialize"),
    )

    with pytest.raises(RendererCleanupError) as raised:
        _powerpoint_export_lifecycle(
            tmp_path / "deck.pptx",
            raw,
            expected_count=1,
            timeout_seconds=1,
            pythoncom_module=pythoncom,
            dispatch_ex=lambda progid: FakeApplication(),
        )

    assert calls == ["CoInitialize", "Export", "Close", "Quit", "CoUninitialize"]
    assert raised.value.failures == ("Close: OSError",)


def test_owned_renderer_worker_timeout_terminates_the_worker() -> None:
    from presentation_studio.backends.pptx_native import (
        RendererTimeoutError,
        _run_owned_renderer_process,
    )

    started = time.monotonic()
    with pytest.raises(RendererTimeoutError) as raised:
        _run_owned_renderer_process(
            _blocking_renderer_worker,
            (10.0,),
            timeout_seconds=0.15,
        )

    assert time.monotonic() - started < 3
    assert raised.value.process_result.timed_out is True
    assert raised.value.process_result.failure_kind == "timeout"
    assert not _pid_is_running(raised.value.worker_pid)


def test_post_start_keyboard_interrupt_preserves_primary_and_stops_real_worker(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    primary_error = KeyboardInterrupt("synthetic interruption after child creation")

    caught, worker_was_alive = _run_supervisor_with_post_start_failure(
        monkeypatch, primary_error
    )

    assert caught is primary_error
    assert worker_was_alive is False


def test_post_start_runtime_error_preserves_primary_and_stops_real_worker(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    primary_error = RuntimeError("synthetic start failure after child creation")

    caught, worker_was_alive = _run_supervisor_with_post_start_failure(
        monkeypatch, primary_error
    )

    assert caught is primary_error
    assert worker_was_alive is False


def test_pre_start_error_preserves_primary_without_worker_cleanup(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from presentation_studio.backends import pptx_native as native

    process_type = type(native.multiprocessing.get_context("spawn").Process())
    primary_error = OSError("synthetic failure before child creation")
    cleanup_calls: list[object] = []

    def fail_before_start(process: object) -> None:
        del process
        raise primary_error

    monkeypatch.setattr(process_type, "start", fail_before_start)
    monkeypatch.setattr(native, "_stop_worker", cleanup_calls.append)

    with pytest.raises(OSError) as raised:
        native._run_owned_renderer_process(
            _blocking_renderer_worker,
            (30.0,),
            timeout_seconds=2,
        )

    assert raised.value is primary_error
    assert cleanup_calls == []


def test_worker_shutdown_kills_and_joins_when_terminate_raises() -> None:
    from presentation_studio.backends import pptx_native as native

    worker = _TerminateRaisesWorker()

    with pytest.raises(native.RendererCleanupError) as raised:
        native._stop_worker(worker)

    assert worker.calls == ["terminate", "join", "kill", "join"]
    assert worker.is_alive() is False
    assert raised.value.failures == ("worker terminate: OSError",)


def test_external_shutdown_kills_and_waits_when_terminate_raises() -> None:
    from presentation_studio.backends import pptx_native as native

    external = _TerminateRaisesExternal()

    with pytest.raises(native.RendererCleanupError) as raised:
        native._stop_owned_external(external)

    assert external.calls == ["terminate", "wait", "kill", "wait"]
    assert external.poll() == -9
    assert raised.value.failures == ("owned renderer terminate: OSError",)


@pytest.mark.skipif(sys.platform != "win32", reason="Windows owned process handle")
def test_owned_renderer_timeout_terminates_published_external_process(
    tmp_path: Path,
) -> None:
    from presentation_studio.backends.pptx_native import (
        RendererTimeoutError,
        _run_owned_renderer_process,
    )

    pid_file = tmp_path / "external.pid"

    def launch_owned_process():
        child = subprocess.Popen(
            [sys.executable, "-c", "import time; time.sleep(30)"],
            stdin=subprocess.DEVNULL,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        pid_file.write_text(str(child.pid), encoding="ascii")
        return child

    with pytest.raises(RendererTimeoutError):
        _run_owned_renderer_process(
            _blocking_renderer_worker,
            (10.0,),
            timeout_seconds=1.5,
            external_launcher=launch_owned_process,
        )

    external_pid = int(pid_file.read_text(encoding="ascii"))
    assert not _pid_is_running(external_pid)


@pytest.mark.skipif(sys.platform != "win32", reason="Windows PowerPoint ownership")
def test_supervisor_never_kills_foreign_powerpoint_from_pid_delta(tmp_path: Path) -> None:
    from presentation_studio.backends import pptx_native as native

    discovery = discover_renderer()
    if not discovery.available or discovery.executable is None:
        pytest.skip(discovery.unavailable_reason or "PowerPoint renderer unavailable")
    baseline = tuple(native._powerpoint_pids())
    foreign = subprocess.Popen(
        [str(discovery.executable), "/automation"],
        stdin=subprocess.DEVNULL,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    worker_pid_file = tmp_path / "worker.pid"
    survived = False
    try:
        time.sleep(0.25)
        with pytest.raises(native.RendererTimeoutError):
            native._run_owned_renderer_process(
                _baseline_only_renderer_worker,
                (baseline, str(worker_pid_file), 10.0),
                timeout_seconds=1.5,
            )
        survived = foreign.poll() is None
    finally:
        if foreign.poll() is None:
            foreign.terminate()
            try:
                foreign.wait(timeout=5)
            except subprocess.TimeoutExpired:
                foreign.kill()
                foreign.wait(timeout=5)

    assert survived, "supervisor killed a foreign PowerPoint inferred from PID delta"
    assert not _pid_is_running(int(worker_pid_file.read_text(encoding="ascii")))


@pytest.mark.skipif(sys.platform != "win32", reason="Windows owned process handle")
def test_keyboard_interrupt_after_bind_stops_worker_and_owned_process(
    tmp_path: Path,
) -> None:
    from presentation_studio.backends import pptx_native as native

    pid_file = tmp_path / "external.pid"
    worker_pid_file = tmp_path / "worker.pid"

    def launch_owned_process():
        child = subprocess.Popen(
            [sys.executable, "-c", "import time; time.sleep(30)"],
            stdin=subprocess.DEVNULL,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        pid_file.write_text(str(child.pid), encoding="ascii")
        return child

    def interrupt_after_both_processes_exist() -> None:
        deadline = time.monotonic() + 5
        while time.monotonic() < deadline:
            if pid_file.exists() and worker_pid_file.exists():
                time.sleep(0.1)
                _thread.interrupt_main()
                return
            time.sleep(0.01)

    interrupter = Thread(target=interrupt_after_both_processes_exist, daemon=True)
    interrupter.start()
    interrupted = False
    worker_alive = True
    external_alive = True
    try:
        with pytest.raises(KeyboardInterrupt):
            native._run_owned_renderer_process(
                _blocking_renderer_worker_with_pid,
                (str(worker_pid_file), 10.0),
                timeout_seconds=5,
                external_launcher=launch_owned_process,
            )
        interrupted = True
        worker_pid = int(worker_pid_file.read_text(encoding="ascii"))
        external_pid = int(pid_file.read_text(encoding="ascii"))
        worker_alive = _pid_is_running(worker_pid)
        external_alive = _pid_is_running(external_pid)
    finally:
        interrupter.join(timeout=1)
        for path in (pid_file, worker_pid_file):
            if not path.exists():
                continue
            pid = int(path.read_text(encoding="ascii"))
            subprocess.run(
                ["taskkill", "/PID", str(pid), "/T", "/F"],
                check=False,
                stdin=subprocess.DEVNULL,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
            )

    assert interrupted
    assert worker_alive is False
    assert external_alive is False


@pytest.mark.skipif(sys.platform != "win32", reason="Windows process liveness")
def test_pid_liveness_helper_distinguishes_live_and_joined_processes() -> None:
    assert _pid_is_running(os.getpid()) is True
    child = subprocess.Popen(
        [sys.executable, "-c", "import time; time.sleep(30)"],
        stdin=subprocess.DEVNULL,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    try:
        assert _pid_is_running(child.pid) is True
    finally:
        child.terminate()
        child.wait(timeout=5)
    assert _pid_is_running(child.pid) is False


@pytest.mark.skipif(sys.platform != "win32", reason="Windows process liveness")
def test_pid_liveness_helper_treats_query_failure_as_unknown_error(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import ctypes

    class FakeFunction:
        def __init__(self, callback) -> None:
            self.callback = callback
            self.argtypes = None
            self.restype = None

        def __call__(self, *args):
            return self.callback(*args)

    class FakeKernel:
        OpenProcess = FakeFunction(lambda access, inherit, pid: 123)
        GetExitCodeProcess = FakeFunction(
            lambda handle, output: (ctypes.set_last_error(5), 0)[1]
        )
        CloseHandle = FakeFunction(lambda handle: 1)

    monkeypatch.setattr(ctypes, "WinDLL", lambda *args, **kwargs: FakeKernel())

    with pytest.raises(OSError):
        _pid_is_running(123)


@pytest.mark.skipif(sys.platform != "win32", reason="Windows owned process cleanup")
@pytest.mark.parametrize(
    ("entrypoint", "expected_error"),
    [
        (_crashing_renderer_worker, "RendererLifecycleError"),
        (_error_renderer_worker, "RendererLifecycleError"),
    ],
)
def test_every_worker_failure_branch_joins_worker_and_owned_process(
    tmp_path: Path, entrypoint, expected_error: str
) -> None:
    from presentation_studio.backends import pptx_native as native

    worker_pid_file = tmp_path / f"{entrypoint.__name__}-worker.pid"
    external_pid_file = tmp_path / f"{entrypoint.__name__}-external.pid"

    def launch_owned_process():
        child = subprocess.Popen(
            [sys.executable, "-c", "import time; time.sleep(30)"],
            stdin=subprocess.DEVNULL,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        external_pid_file.write_text(str(child.pid), encoding="ascii")
        return child

    with pytest.raises(getattr(native, expected_error)):
        native._run_owned_renderer_process(
            entrypoint,
            (str(worker_pid_file),),
            timeout_seconds=3,
            external_launcher=launch_owned_process,
        )

    assert not _pid_is_running(int(worker_pid_file.read_text(encoding="ascii")))
    assert not _pid_is_running(int(external_pid_file.read_text(encoding="ascii")))


@pytest.mark.skipif(sys.platform != "win32", reason="Windows PowerPoint ownership")
def test_real_renderer_fails_closed_when_foreign_powerpoint_exists(
    tmp_path: Path,
) -> None:
    from presentation_studio.backends import pptx_native as native

    discovery = discover_renderer()
    if not discovery.available or discovery.executable is None:
        pytest.skip(discovery.unavailable_reason or "PowerPoint renderer unavailable")
    paths = resolve_paths(workspace=tmp_path)
    build = _backend_without_renderer().build(_build_input(_deck()), paths)
    foreign = subprocess.Popen(
        [str(discovery.executable), "/automation"],
        stdin=subprocess.DEVNULL,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    survived = False
    try:
        time.sleep(0.25)
        with pytest.raises(native.RendererLifecycleError) as raised:
            PptxNativeBackend(renderer=discovery).render(
                paths, build.build_id, timeout_seconds=30
            )
        survived = foreign.poll() is None
    finally:
        if foreign.poll() is None:
            foreign.terminate()
            try:
                foreign.wait(timeout=5)
            except subprocess.TimeoutExpired:
                foreign.kill()
                foreign.wait(timeout=5)

    assert raised.value.primary_type == "RendererOwnershipUnproven"
    assert survived is True


@pytest.mark.skipif(sys.platform != "win32", reason="Windows PowerPoint ownership")
def test_real_renderer_does_not_kill_foreign_powerpoint_started_during_dispatch(
    tmp_path: Path,
) -> None:
    from presentation_studio.backends import pptx_native as native

    discovery = discover_renderer()
    if not discovery.available or discovery.executable is None:
        pytest.skip(discovery.unavailable_reason or "PowerPoint renderer unavailable")
    paths = resolve_paths(workspace=tmp_path)
    build = _backend_without_renderer().build(_build_input(_deck()), paths)
    foreign_holder: list[subprocess.Popen] = []

    def launch_foreign_after_owned_appears() -> None:
        deadline = time.monotonic() + 5
        while time.monotonic() < deadline:
            if native._powerpoint_pids():
                foreign_holder.append(
                    subprocess.Popen(
                        [str(discovery.executable), "/automation"],
                        stdin=subprocess.DEVNULL,
                        stdout=subprocess.DEVNULL,
                        stderr=subprocess.DEVNULL,
                    )
                )
                return
            time.sleep(0.005)

    starter = Thread(target=launch_foreign_after_owned_appears, daemon=True)
    starter.start()
    error: BaseException | None = None
    survived = False
    try:
        with pytest.raises(native.RendererLifecycleError) as raised:
            PptxNativeBackend(renderer=discovery).render(
                paths, build.build_id, timeout_seconds=30
            )
        error = raised.value
        starter.join(timeout=5)
        assert foreign_holder, "foreign PowerPoint did not start during renderer lifecycle"
        survived = foreign_holder[0].poll() is None
    finally:
        starter.join(timeout=1)
        for foreign in foreign_holder:
            if foreign.poll() is None:
                foreign.terminate()
                try:
                    foreign.wait(timeout=5)
                except subprocess.TimeoutExpired:
                    foreign.kill()
                    foreign.wait(timeout=5)

    assert isinstance(error, native.RendererLifecycleError)
    assert survived is True


@pytest.mark.skipif(sys.platform != "win32", reason="Windows PowerPoint concurrency")
def test_concurrent_real_renders_fail_closed_without_cross_kill_or_orphan(
    tmp_path: Path,
) -> None:
    from presentation_studio.backends import pptx_native as native

    discovery = discover_renderer()
    if not discovery.available:
        pytest.skip(discovery.unavailable_reason or "PowerPoint renderer unavailable")
    paths = resolve_paths(workspace=tmp_path)
    builds = [
        _backend_without_renderer().build(_build_input(_deck()), paths)
        for _ in range(2)
    ]
    barrier = Barrier(2)
    outcomes: list[object] = []

    def render(build_id: str) -> None:
        barrier.wait(timeout=5)
        try:
            outcomes.append(
                PptxNativeBackend(renderer=discovery).render(
                    paths, build_id, timeout_seconds=30
                )
            )
        except BaseException as exc:
            outcomes.append(exc)

    threads = [Thread(target=render, args=(build.build_id,)) for build in builds]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join(timeout=20)

    assert all(not thread.is_alive() for thread in threads)
    assert len(outcomes) == 2
    assert all(
        getattr(outcome, "status", None) == "passed"
        or isinstance(outcome, native.RendererLifecycleError)
        for outcome in outcomes
    )
    assert not native._powerpoint_pids()


def test_project_render_timeout_is_forwarded_to_the_whole_renderer_lifecycle(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from presentation_studio.backends import pptx_native as native

    paths = resolve_paths(workspace=tmp_path)
    config = ProjectConfig(
        project_id="timeout-demo",
        title="Timeout demo",
        limits={"render_timeout_seconds": 0.25},
    )
    (tmp_path / "project.yaml").write_text(
        yaml.safe_dump(config.model_dump(mode="json"), allow_unicode=True),
        encoding="utf-8",
    )
    build_id = "timeout-build"
    source = tmp_path / "builds" / build_id / "deck.pptx"
    source.parent.mkdir(parents=True)
    source.write_bytes(b"fixture")
    monkeypatch.setattr(
        native,
        "extract_semantics",
        lambda path: [SimpleNamespace(slide_id="s01")],
    )
    observed: list[float] = []

    def export_runner(source_path, raw_directory, expected_count, timeout_seconds):
        del source_path, expected_count
        observed.append(timeout_seconds)
        Image.new("RGB", (40, 30), "green").save(raw_directory / "Slide1.PNG")

    backend = PptxNativeBackend(
        renderer=RendererDiscovery(renderer_id="powerpoint-com", available=True),
        renderer_runner=export_runner,
    )

    backend.render(paths, build_id)

    assert observed == [0.25]


def test_post_promotion_hash_validation_rejects_tampered_render(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from presentation_studio.backends import pptx_native as native
    from presentation_studio.fs import BoundDirectory

    paths = resolve_paths(workspace=tmp_path)
    build_id = "tamper-build"
    source = tmp_path / "builds" / build_id / "deck.pptx"
    source.parent.mkdir(parents=True)
    source.write_bytes(b"fixture")
    monkeypatch.setattr(
        native,
        "extract_semantics",
        lambda path: [SimpleNamespace(slide_id="s01")],
    )

    def export_runner(source_path, raw_directory, expected_count, timeout_seconds):
        del source_path, expected_count, timeout_seconds
        Image.new("RGB", (40, 30), "green").save(raw_directory / "Slide1.PNG")

    original_promote = BoundDirectory.rename_noreplace

    def tamper_after_promotion(binding: BoundDirectory, target: Path) -> None:
        original_promote(binding, target)
        if target.name.startswith("render-"):
            next(target.glob("*.png")).write_bytes(b"tampered")

    monkeypatch.setattr(BoundDirectory, "rename_noreplace", tamper_after_promotion)
    backend = PptxNativeBackend(
        renderer=RendererDiscovery(renderer_id="powerpoint-com", available=True),
        renderer_runner=export_runner,
    )

    with pytest.raises(native.RenderArtifactError):
        backend.render(paths, build_id, timeout_seconds=1)

    render_root = tmp_path / "builds" / build_id / "renders"
    assert not list(render_root.glob("render-*"))


@pytest.mark.parametrize("chart_type", ["bar", "column", "line", "pie"])
def test_all_supported_chart_types_round_trip_natively(
    tmp_path: Path, chart_type: str
) -> None:
    deck = DeckSpec(
        title="Biểu đồ native",
        audience="A",
        purpose="P",
        slides=[
            {
                "slide_id": "s01",
                "title": chart_type,
                "layout_ref": "builtin:data@1.0.0",
                "elements": [
                    {
                        "element_id": "e-chart",
                        "kind": "chart",
                        "content": {
                            "chart_type": chart_type,
                            "data": [{"label": "A", "value": 1}, {"label": "B", "value": 2}],
                            "unit": "điểm",
                        },
                    }
                ],
            }
        ],
    )
    paths = resolve_paths(workspace=tmp_path)

    result = _backend_without_renderer().build(_build_input(deck), paths)
    semantic = extract_semantics(tmp_path / result.outputs[0].path)

    assert semantic[0].elements["e-chart"].chart_type == chart_type


def test_all_six_supported_element_kinds_round_trip_natively(tmp_path: Path) -> None:
    elements = [
        {"element_id": "e-text", "kind": "text", "content": "Nội dung"},
        {
            "element_id": "e-table",
            "kind": "table",
            "content": {"columns": ["A"], "rows": [[1]]},
        },
        {
            "element_id": "e-chart",
            "kind": "chart",
            "content": {"chart_type": "bar", "data": [{"label": "A", "value": 1}]},
        },
        {
            "element_id": "e-process",
            "kind": "process",
            "content": {"steps": [{"id": "one", "title": "Bước một"}]},
        },
        {
            "element_id": "e-timeline",
            "kind": "timeline",
            "content": {"items": [{"time_label": "2026", "title": "Mốc một"}]},
        },
        {
            "element_id": "e-quote",
            "kind": "quote",
            "content": {"quote": "Trích dẫn", "attribution": "Tác giả"},
        },
    ]
    slides = [
        {
            "slide_id": f"s{index:02d}",
            "title": element["kind"],
            "layout_ref": "builtin:content@1.0.0",
            "elements": [element],
        }
        for index, element in enumerate(elements, start=1)
    ]
    deck = DeckSpec(title="Sáu loại", audience="A", purpose="P", slides=slides)
    paths = resolve_paths(workspace=tmp_path)

    result = _backend_without_renderer().build(_build_input(deck), paths)
    semantic = extract_semantics(tmp_path / result.outputs[0].path)

    assert [next(iter(slide.elements.values())).kind for slide in semantic] == [
        "text",
        "table",
        "chart",
        "process",
        "timeline",
        "quote",
    ]


def test_semantic_corruption_and_reopen_failure_never_promote_build(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from presentation_studio.backends import pptx_native as native

    paths = resolve_paths(workspace=tmp_path)
    monkeypatch.setattr(native, "extract_semantics", lambda path: [])

    with pytest.raises(NativeBuildError):
        _backend_without_renderer().build(_build_input(_deck()), paths)

    assert not [path for path in paths.builds_dir.glob("native-*")]


def test_reopen_exception_never_promotes_build(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from presentation_studio.backends import pptx_native as native

    paths = resolve_paths(workspace=tmp_path)

    def fail_reopen(path: Path):
        raise OSError("synthetic reopen failure")

    monkeypatch.setattr(native, "extract_semantics", fail_reopen)

    with pytest.raises(OSError, match="synthetic reopen failure"):
        _backend_without_renderer().build(_build_input(_deck()), paths)

    assert not [path for path in paths.builds_dir.glob("native-*")]


def test_save_failure_never_promotes_build(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from pptx.presentation import Presentation as PresentationClass

    paths = resolve_paths(workspace=tmp_path)

    def fail_save(self, path) -> None:
        raise OSError("synthetic save failure")

    monkeypatch.setattr(PresentationClass, "save", fail_save)

    with pytest.raises(OSError, match="synthetic save failure"):
        _backend_without_renderer().build(_build_input(_deck()), paths)

    assert not [path for path in paths.builds_dir.glob("native-*")]


def test_edit_then_real_powerpoint_render_preserves_edit_and_renders_all_slides(
    tmp_path: Path,
) -> None:
    discovery = discover_renderer()
    if not discovery.available:
        pytest.skip(discovery.unavailable_reason or "PowerPoint renderer unavailable")
    paths = resolve_paths(workspace=tmp_path)
    backend = PptxNativeBackend(renderer=discovery)
    build = backend.build(_build_input(_deck()), paths)
    output = tmp_path / build.outputs[0].path
    presentation = Presentation(output)
    edited = next(
        shape
        for shape in presentation.slides[0].shapes
        if shape.name == "presentation-studio:element:e-explain:text"
    )
    edited.text = "Nội dung đã sửa trước khi render."
    presentation.save(output)

    result = backend.render(paths, build.build_id, timeout_seconds=30)

    assert result.status == "passed"
    assert list(result.slides) == build.expected_slide_ids
    assert all((tmp_path / slide.path).is_file() for slide in result.slides.values())
