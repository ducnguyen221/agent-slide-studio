from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime, timezone
import hashlib
import importlib.util
import io
import json
import multiprocessing
import os
from pathlib import Path
import re
import shutil
import stat
import subprocess
import sys
import time
from uuid import uuid4

import pptx
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

from presentation_studio.fs import BoundDirectory, UnsafeFileError
from presentation_studio.layout.native import (
    CHART_DETAIL_FONT_PT,
    MESSAGE_FONT_PT,
    METRIC_FONT_PT,
    QUOTE_FONT_PT,
    SHAPE_FONT_PT,
    SOURCE_FONT_PT,
    TABLE_BODY_FONT_PT,
    TABLE_HEADER_FONT_PT,
    TEXT_FONT_PT,
    TITLE_FONT_PT,
    LayoutIssue,
    NativeBox,
    canvas_inches,
    content_boxes,
    message_box,
    preflight_deck,
    source_box,
    title_box,
)
from presentation_studio.models import (
    BackendCapabilities,
    BuildInput,
    BuildResult,
    ChartElement,
    CLIError,
    CLIResult,
    FeatureVerification,
    OutputArtifact,
    ProcessElement,
    ProcessResult,
    QuoteElement,
    RenderCapabilities,
    RenderDimensions,
    RenderedSlide,
    RenderResult,
    SlideBuildResult,
    SlideSpec,
    TableElement,
    TextContent,
    TextElement,
    TimelineElement,
)
from presentation_studio.state import hash_inputs
from presentation_studio.workspace import (
    WorkspacePaths,
    load_deck,
    load_project,
    resolve_profile,
    safe_path,
)

from .base import PresentationBackend


PPTX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.presentationml.presentation"
)
_ELEMENT_NAME = re.compile(
    r"^presentation-studio:element:(?P<element_id>[a-z0-9][a-z0-9-]{0,63}):(?P<kind>[a-z0-9-]+)$"
)
_SLIDE_NAME = re.compile(
    r"^presentation-studio:slide:(?P<slide_id>[a-z0-9][a-z0-9-]{0,63}):(?P<field>title|message|sources)$"
)
_IDENTIFIER = re.compile(r"^[a-z0-9][a-z0-9-]{0,63}$")


class RenderArtifactError(RuntimeError):
    pass


class RendererCleanupError(RuntimeError):
    def __init__(self, failures: list[str] | tuple[str, ...]):
        self.failures = tuple(failures)
        super().__init__("PowerPoint cleanup failed: " + ", ".join(self.failures))


class RendererLifecycleError(RuntimeError):
    def __init__(
        self,
        primary_type: str,
        primary_message: str,
        *,
        notes: list[str] | tuple[str, ...] = (),
        process_result: ProcessResult | None = None,
    ) -> None:
        self.primary_type = primary_type
        self.primary_message = primary_message
        self.process_result = process_result
        super().__init__(f"{primary_type}: {primary_message}")
        for note in notes:
            self.add_note(note)


class RendererTimeoutError(TimeoutError):
    def __init__(self, worker_pid: int, process_result: ProcessResult):
        self.worker_pid = worker_pid
        self.process_result = process_result
        super().__init__(
            f"Renderer vượt quá deadline {process_result.elapsed_seconds:.3f} giây."
        )


@dataclass(frozen=True)
class RendererDiscovery:
    renderer_id: str
    available: bool
    unavailable_reason: str | None = None
    executable: Path | None = None
    version: str | None = None


@dataclass(frozen=True)
class SemanticElement:
    kind: str
    text: str | None = None
    columns: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    chart_type: str | None = None
    data: list[tuple[str, float]] = field(default_factory=list)
    unit: str | None = None
    source_ref: str | None = None


@dataclass(frozen=True)
class SlideSemantic:
    slide_id: str
    title: str
    message: str | None
    notes: str | None
    source_refs: list[str]
    elements: dict[str, SemanticElement]


@dataclass
class _BoundStage:
    """Lifecycle wrapper over WP02's handle-bound filesystem primitive."""

    path: Path
    parent: BoundDirectory
    directory: BoundDirectory
    device: int
    inode: int
    removed: bool = False

    @classmethod
    def create(cls, parent_path: Path, name: str) -> _BoundStage:
        parent = BoundDirectory.open(parent_path)
        directory: BoundDirectory | None = None
        try:
            directory = parent.child(
                name,
                create=True,
                exclusive=True,
                allow_delete=True,
            )
            info = os.fstat(directory.descriptor)
            return cls(
                path=directory.path,
                parent=parent,
                directory=directory,
                device=info.st_dev,
                inode=info.st_ino,
            )
        except BaseException as primary_error:
            failures: list[str] = []
            for label, binding in (("directory", directory), ("parent", parent)):
                if binding is None:
                    continue
                try:
                    binding.close()
                except BaseException as cleanup_error:
                    failures.append(f"{label}: {type(cleanup_error).__name__}")
            if failures:
                primary_error.add_note("staging capture cleanup failed: " + ", ".join(failures))
            raise

    def verify(self) -> None:
        if self.removed:
            raise UnsafeFileError("staging directory was already removed")
        self.parent.verify()
        self.directory.verify()
        current = self.parent.lstat(self.path.name)
        if (current.st_dev, current.st_ino) != (self.device, self.inode):
            raise UnsafeFileError("staging directory identity changed")

    def promote(self, target: Path) -> None:
        self.verify()
        if target.parent.absolute() != self.parent.path.absolute():
            raise UnsafeFileError("staging target parent does not match bound parent")
        self.directory.rename_noreplace(target)
        self.path = target.absolute()
        self.verify()

    def remove(self) -> None:
        self.verify()
        if os.name == "nt":
            self.directory.remove_tree()
        else:
            quarantine = f"cleanup-{uuid4().hex}"
            self.parent.replace(self.path.name, quarantine)
            moved = self.parent.lstat(quarantine)
            if (moved.st_dev, moved.st_ino) != (self.device, self.inode):
                self.parent.replace(quarantine, self.path.name)
                raise UnsafeFileError("staging directory identity changed during cleanup")
            shutil.rmtree(quarantine, dir_fd=self.parent.descriptor)
            self.parent.fsync()
        self.removed = True

    def close(self) -> None:
        failures: list[BaseException] = []
        for binding in (self.directory, self.parent):
            try:
                binding.close()
            except BaseException as error:
                failures.append(error)
        if failures:
            raise failures[0]


class NativeBuildError(ValueError):
    def __init__(self, issues: list[LayoutIssue]):
        self.issues = tuple(issues)
        super().__init__("; ".join(issue.message_vi for issue in issues))


def _file_version(path: Path) -> str | None:
    try:
        import win32api

        info = win32api.GetFileVersionInfo(str(path), "\\")
        major = info["FileVersionMS"] >> 16
        minor = info["FileVersionMS"] & 0xFFFF
        build = info["FileVersionLS"] >> 16
        revision = info["FileVersionLS"] & 0xFFFF
        return f"{major}.{minor}.{build}.{revision}"
    except (ImportError, OSError, KeyError, TypeError):
        return None


def discover_renderer() -> RendererDiscovery:
    """Phát hiện renderer cục bộ mà không khởi chạy ứng dụng."""
    if sys.platform != "win32":
        return RendererDiscovery(
            renderer_id="none",
            available=False,
            unavailable_reason="Backend native hiện chỉ tích hợp renderer PowerPoint COM trên Windows.",
        )
    try:
        win32com_spec = importlib.util.find_spec("win32com.client")
    except ModuleNotFoundError:
        win32com_spec = None
    if win32com_spec is None:
        return RendererDiscovery(
            renderer_id="none",
            available=False,
            unavailable_reason="Thiếu pywin32 để điều khiển PowerPoint renderer.",
        )
    try:
        import winreg

        locations = (
            (winreg.HKEY_LOCAL_MACHINE, winreg.KEY_READ | winreg.KEY_WOW64_64KEY),
            (winreg.HKEY_LOCAL_MACHINE, winreg.KEY_READ | winreg.KEY_WOW64_32KEY),
            (winreg.HKEY_CURRENT_USER, winreg.KEY_READ),
        )
        for hive, access in locations:
            try:
                with winreg.OpenKey(
                    hive,
                    r"SOFTWARE\Microsoft\Windows\CurrentVersion\App Paths\POWERPNT.EXE",
                    0,
                    access,
                ) as key:
                    executable = Path(winreg.QueryValue(key, None)).resolve()
                if executable.is_file():
                    return RendererDiscovery(
                        renderer_id="powerpoint-com",
                        available=True,
                        executable=executable,
                        version=_file_version(executable),
                    )
            except (FileNotFoundError, OSError):
                continue
    except ImportError:
        pass
    return RendererDiscovery(
        renderer_id="none",
        available=False,
        unavailable_reason="Không tìm thấy Microsoft PowerPoint renderer trong registry.",
    )


def native_capabilities(
    renderer: RendererDiscovery | None = None,
) -> BackendCapabilities:
    renderer = renderer or discover_renderer()
    promotion_available = _identity_safe_native_promotion_available()
    visual_reason = renderer.unavailable_reason or (
        "Renderer đã được phát hiện nhưng chưa có evidence của lượt render này."
    )
    return BackendCapabilities(
        backend_id="pptx-native",
        adapter_version="1.0.0",
        engine_version=pptx.__version__,
        available=promotion_available,
        unavailable_reason=(
            None
            if promotion_available
            else "Chưa có identity-safe promotion đã kiểm chứng trên nền tảng này."
        ),
        supported_inputs=["semantic"],
        supported_outputs=["pptx"],
        supported_elements=["chart", "process", "quote", "table", "text", "timeline"],
        editable_elements=["chart", "process", "quote", "table", "text", "timeline"],
        notes_support_by_output={"pptx": True},
        animation_support=False,
        requires_network=False,
        render_capabilities=RenderCapabilities(
            png=renderer.available,
            pdf=False,
            dimensions=renderer.available,
        ),
        verification_by_feature={
            "semantic-reopen": FeatureVerification(
                status="unverified",
                environment=f"python-pptx {pptx.__version__}",
            ),
            "visual-render": FeatureVerification(
                status="unverified",
                environment=visual_reason,
            ),
        },
    )


def _identity_safe_native_promotion_available() -> bool:
    """Return true only where promotion is pinned to an owned directory handle."""

    return os.name == "nt"


def _metadata_node(shape: object) -> object:
    nodes = shape._element.xpath(".//p:cNvPr")
    if not nodes:
        raise ValueError("shape does not expose cNvPr metadata")
    return nodes[0]


def _tag_shape(shape: object, name: str, metadata: dict[str, object]) -> None:
    shape.name = name
    node = _metadata_node(shape)
    node.set("title", json.dumps(metadata, ensure_ascii=False, separators=(",", ":")))
    alt_text = metadata.get("alt_text")
    node.set("descr", alt_text if isinstance(alt_text, str) else "")


def _shape_metadata(shape: object) -> dict[str, object]:
    raw = _metadata_node(shape).get("title")
    if not raw:
        return {}
    try:
        value = json.loads(raw)
    except json.JSONDecodeError:
        return {}
    return value if isinstance(value, dict) else {}


def _format_text(element: TextElement) -> str:
    if isinstance(element.content, str):
        return element.content
    content: TextContent = element.content
    parts = [part for part in (content.label, content.text) if part]
    parts.extend(content.items)
    return "\n".join(parts)


def _format_process(element: ProcessElement) -> str:
    return "\n".join(
        f"{index}. {step.title}" + (f" — {step.description}" if step.description else "")
        for index, step in enumerate(element.content.steps, start=1)
    )


def _format_timeline(element: TimelineElement) -> str:
    return "\n".join(
        f"{item.time_label}: {item.title}" + (f" — {item.description}" if item.description else "")
        for item in element.content.items
    )


def _format_quote(element: QuoteElement) -> str:
    text = f"“{element.content.quote}”"
    if element.content.attribution:
        text += f"\n— {element.content.attribution}"
    return text


def _set_text_style(
    shape: object,
    *,
    size: float,
    bold: bool = False,
    horizontal_margin: float = 0.16,
    vertical_margin: float = 0.10,
) -> None:
    frame = shape.text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = Inches(horizontal_margin)
    frame.margin_right = Inches(horizontal_margin)
    frame.margin_top = Inches(vertical_margin)
    frame.margin_bottom = Inches(vertical_margin)
    for paragraph in frame.paragraphs:
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = RGBColor(31, 41, 55)


def _add_textbox(
    slide: object,
    *,
    left: float,
    top: float,
    width: float,
    height: float,
    text: str,
    name: str,
    metadata: dict[str, object],
    size: float,
    bold: bool = False,
    horizontal_margin: float = 0.16,
    vertical_margin: float = 0.10,
) -> object:
    shape = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.text = text
    _set_text_style(
        shape,
        size=size,
        bold=bold,
        horizontal_margin=horizontal_margin,
        vertical_margin=vertical_margin,
    )
    _tag_shape(shape, name, metadata)
    return shape


def _render_text(slide: object, element: TextElement, box: NativeBox) -> None:
    left, top, width, height = box.left, box.top, box.width, box.height
    is_metric = element.role == "metric"
    if is_metric:
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(left),
            Inches(top),
            Inches(width),
            Inches(height),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(239, 246, 255)
        shape.line.color.rgb = RGBColor(37, 99, 235)
        shape.text = _format_text(element)
        _set_text_style(shape, size=METRIC_FONT_PT, bold=True)
        shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        _tag_shape(
            shape,
            f"presentation-studio:element:{element.element_id}:text",
            {"kind": "text", "role": element.role, "alt_text": element.alt_text},
        )
        return
    _add_textbox(
        slide,
        left=left,
        top=top,
        width=width,
        height=height,
        text=_format_text(element),
        name=f"presentation-studio:element:{element.element_id}:text",
        metadata={"kind": "text", "role": element.role, "alt_text": element.alt_text},
        size=TEXT_FONT_PT,
    )


def _render_table(slide: object, element: TableElement, box: NativeBox) -> None:
    left, top, width, height = box.left, box.top, box.width, box.height
    rows = len(element.content.rows) + 1
    columns = len(element.content.columns)
    shape = slide.shapes.add_table(
        rows,
        columns,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    table = shape.table
    values = [element.content.columns, *element.content.rows]
    for row_index, row in enumerate(values):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.text = "" if value is None else str(value)
            cell.text_frame.word_wrap = True
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name = "Aptos"
                paragraph.font.size = Pt(TABLE_BODY_FONT_PT if row_index else TABLE_HEADER_FONT_PT)
                paragraph.font.bold = row_index == 0
    _tag_shape(
        shape,
        f"presentation-studio:element:{element.element_id}:table",
        {"kind": "table", "role": element.role, "alt_text": element.alt_text},
    )


_CHART_TYPES = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
}


def _render_chart(slide: object, element: ChartElement, box: NativeBox) -> None:
    left, top, width, height = box.left, box.top, box.width, box.height
    data = CategoryChartData()
    data.categories = [datum.label for datum in element.content.data]
    data.add_series("Giá trị", [datum.value for datum in element.content.data])
    shape = slide.shapes.add_chart(
        _CHART_TYPES[element.content.chart_type],
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height - 0.34),
        data,
    )
    shape.chart.has_legend = False
    _tag_shape(
        shape,
        f"presentation-studio:element:{element.element_id}:chart",
        {
            "kind": "chart",
            "chart_type": element.content.chart_type,
            "unit": element.content.unit,
            "source_ref": element.content.source_ref,
            "role": element.role,
            "alt_text": element.alt_text,
        },
    )
    details = []
    if element.content.unit:
        details.append(f"Đơn vị: {element.content.unit}")
    if element.content.source_ref:
        details.append(f"Nguồn: {element.content.source_ref}")
    if details:
        _add_textbox(
            slide,
            left=left,
            top=top + height - 0.30,
            width=width,
            height=0.28,
            text=" · ".join(details),
            name=f"presentation-studio:aux:{element.element_id}:chart-details",
            metadata={"auxiliary_for": element.element_id},
            size=CHART_DETAIL_FONT_PT,
            horizontal_margin=0.02,
            vertical_margin=0.01,
        )


def _render_shape_text(
    slide: object,
    element: ProcessElement | TimelineElement | QuoteElement,
    box: NativeBox,
) -> None:
    left, top, width, height = box.left, box.top, box.width, box.height
    if isinstance(element, ProcessElement):
        text = _format_process(element)
        kind = "process"
    elif isinstance(element, TimelineElement):
        text = _format_timeline(element)
        kind = "timeline"
    else:
        text = _format_quote(element)
        kind = "quote"
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(248, 250, 252)
    shape.line.color.rgb = RGBColor(148, 163, 184)
    shape.text = text
    _set_text_style(
        shape,
        size=SHAPE_FONT_PT if kind != "quote" else QUOTE_FONT_PT,
    )
    _tag_shape(
        shape,
        f"presentation-studio:element:{element.element_id}:{kind}",
        {
            "kind": kind,
            "role": element.role,
            "alt_text": element.alt_text,
            "connector": getattr(element.content, "connector", None),
        },
    )


def _render_slide(
    presentation: Presentation,
    slide_spec: SlideSpec,
    canvas_width: float,
    canvas_height: float,
) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = RGBColor(255, 255, 255)

    title_region = title_box(canvas_width, canvas_height)
    _add_textbox(
        slide,
        left=title_region.left,
        top=title_region.top,
        width=title_region.width,
        height=title_region.height,
        text=slide_spec.title,
        name=f"presentation-studio:slide:{slide_spec.slide_id}:title",
        metadata={"slide_id": slide_spec.slide_id, "field": "title"},
        size=TITLE_FONT_PT,
        bold=True,
        vertical_margin=0.05,
    )
    if slide_spec.message:
        message_region = message_box(canvas_width, canvas_height)
        _add_textbox(
            slide,
            left=message_region.left,
            top=message_region.top,
            width=message_region.width,
            height=message_region.height,
            text=slide_spec.message,
            name=f"presentation-studio:slide:{slide_spec.slide_id}:message",
            metadata={"slide_id": slide_spec.slide_id, "field": "message"},
            size=MESSAGE_FONT_PT,
            vertical_margin=0.01,
        )

    for element, box in zip(
        slide_spec.elements,
        content_boxes(slide_spec, canvas_width, canvas_height),
    ):
        if isinstance(element, TextElement):
            _render_text(slide, element, box)
        elif isinstance(element, TableElement):
            _render_table(slide, element, box)
        elif isinstance(element, ChartElement):
            _render_chart(slide, element, box)
        elif isinstance(element, (ProcessElement, TimelineElement, QuoteElement)):
            _render_shape_text(slide, element, box)
        else:
            raise AssertionError(f"preflight allowed unsupported element: {element.kind}")

    source_text = "Nguồn: " + ", ".join(slide_spec.source_refs) if slide_spec.source_refs else ""
    source_region = source_box(canvas_width, canvas_height)
    source_shape = _add_textbox(
        slide,
        left=source_region.left,
        top=source_region.top,
        width=source_region.width,
        height=source_region.height,
        text=source_text,
        name=f"presentation-studio:slide:{slide_spec.slide_id}:sources",
        metadata={"slide_id": slide_spec.slide_id, "source_refs": slide_spec.source_refs},
        size=SOURCE_FONT_PT,
        vertical_margin=0.01,
    )
    source_shape.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    notes_frame = slide.notes_slide.notes_text_frame
    if notes_frame is None:
        raise RuntimeError("PowerPoint package không tạo được vùng speaker notes.")
    notes_frame.text = slide_spec.notes or ""


def extract_semantics(path: object) -> list[SlideSemantic]:
    """Mở lại PPTX và đọc nội dung từ các đối tượng native thật."""
    presentation = Presentation(path)
    result: list[SlideSemantic] = []
    for slide in presentation.slides:
        slide_id: str | None = None
        title = ""
        message: str | None = None
        source_refs: list[str] = []
        elements: dict[str, SemanticElement] = {}
        for shape in slide.shapes:
            slide_match = _SLIDE_NAME.fullmatch(shape.name)
            if slide_match:
                candidate_id = slide_match.group("slide_id")
                slide_id = slide_id or candidate_id
                if slide_id != candidate_id:
                    raise ValueError("slide metadata contains conflicting slide ids")
                field_name = slide_match.group("field")
                if field_name == "title":
                    title = shape.text
                elif field_name == "message":
                    message = shape.text
                else:
                    raw_refs = _shape_metadata(shape).get("source_refs", [])
                    if isinstance(raw_refs, list) and all(isinstance(item, str) for item in raw_refs):
                        source_refs = raw_refs
                continue
            element_match = _ELEMENT_NAME.fullmatch(shape.name)
            if not element_match:
                continue
            element_id = element_match.group("element_id")
            kind = element_match.group("kind")
            metadata = _shape_metadata(shape)
            if kind == "table" and shape.has_table:
                rows = [
                    [cell.text for cell in row.cells]
                    for row in shape.table.rows
                ]
                elements[element_id] = SemanticElement(
                    kind=kind,
                    columns=rows[0] if rows else [],
                    rows=rows[1:],
                )
            elif kind == "chart" and shape.has_chart:
                plot = shape.chart.plots[0]
                categories = [category.label for category in plot.categories]
                values = [float(value) for value in plot.series[0].values]
                elements[element_id] = SemanticElement(
                    kind=kind,
                    chart_type=str(metadata.get("chart_type") or ""),
                    data=list(zip(categories, values)),
                    unit=metadata.get("unit") if isinstance(metadata.get("unit"), str) else None,
                    source_ref=(
                        metadata.get("source_ref")
                        if isinstance(metadata.get("source_ref"), str)
                        else None
                    ),
                )
            elif shape.has_text_frame:
                elements[element_id] = SemanticElement(kind=kind, text=shape.text)
        if slide_id is None:
            raise ValueError("PPTX thiếu slide_id canonical.")
        notes = slide.notes_slide.notes_text_frame
        result.append(
            SlideSemantic(
                slide_id=slide_id,
                title=title,
                message=message,
                notes=notes.text if notes is not None and notes.text else None,
                source_refs=source_refs,
                elements=elements,
            )
        )
    return result


def _powerpoint_pids() -> set[int]:
    import ctypes

    psapi = ctypes.WinDLL("psapi", use_last_error=True)
    kernel = ctypes.WinDLL("kernel32", use_last_error=True)
    psapi.EnumProcesses.argtypes = [
        ctypes.POINTER(ctypes.c_uint32),
        ctypes.c_uint32,
        ctypes.POINTER(ctypes.c_uint32),
    ]
    psapi.EnumProcesses.restype = ctypes.c_int
    kernel.OpenProcess.argtypes = [ctypes.c_uint32, ctypes.c_int, ctypes.c_uint32]
    kernel.OpenProcess.restype = ctypes.c_void_p
    kernel.QueryFullProcessImageNameW.argtypes = [
        ctypes.c_void_p,
        ctypes.c_uint32,
        ctypes.c_wchar_p,
        ctypes.POINTER(ctypes.c_uint32),
    ]
    kernel.QueryFullProcessImageNameW.restype = ctypes.c_int
    kernel.CloseHandle.argtypes = [ctypes.c_void_p]
    kernel.CloseHandle.restype = ctypes.c_int
    process_ids = (ctypes.c_uint32 * 4096)()
    returned = ctypes.c_uint32()
    if not psapi.EnumProcesses(
        process_ids, ctypes.sizeof(process_ids), ctypes.byref(returned)
    ):
        raise ctypes.WinError(ctypes.get_last_error())
    result: set[int] = set()
    for pid in process_ids[: returned.value // ctypes.sizeof(ctypes.c_uint32)]:
        handle = kernel.OpenProcess(0x1000, False, pid)
        if not handle:
            continue
        try:
            capacity = ctypes.c_uint32(32768)
            buffer = ctypes.create_unicode_buffer(capacity.value)
            if kernel.QueryFullProcessImageNameW(
                handle, 0, buffer, ctypes.byref(capacity)
            ) and Path(buffer.value).name.upper() == "POWERPNT.EXE":
                result.add(int(pid))
        finally:
            kernel.CloseHandle(handle)
    return result


def _powerpoint_export_lifecycle(
    source: Path,
    raw_directory: Path,
    *,
    expected_count: int,
    timeout_seconds: float,
    pythoncom_module: object | None = None,
    dispatch_ex: object | None = None,
    get_object: object | None = None,
) -> None:
    """Run every COM phase while preserving primary and cleanup failures."""
    if pythoncom_module is None or (dispatch_ex is None and get_object is None):
        import pythoncom
        from win32com.client import DispatchEx, GetObject

        pythoncom_module = pythoncom
        dispatch_ex = dispatch_ex or DispatchEx
        get_object = get_object or GetObject

    initialized = False
    application = None
    presentation = None
    primary_error: BaseException | None = None
    try:
        pythoncom_module.CoInitialize()
        initialized = True
        if get_object is not None:
            presentation = get_object(str(source.resolve()))
            application = presentation.Application
        else:
            application = dispatch_ex("PowerPoint.Application")
            presentation = application.Presentations.Open(
                str(source.resolve()), ReadOnly=True, Untitled=False, WithWindow=False
            )
        presentation.Export(str(raw_directory.resolve()), "PNG")
        wait_for_exported_slides(
            raw_directory,
            expected_count=expected_count,
            timeout_seconds=timeout_seconds,
        )
    except BaseException as error:
        primary_error = error

    cleanup_failures: list[str] = []
    for label, action in (
        ("Close", presentation.Close if presentation is not None else None),
        ("Quit", application.Quit if application is not None else None),
        (
            "CoUninitialize",
            pythoncom_module.CoUninitialize if initialized else None,
        ),
    ):
        if action is None:
            continue
        try:
            action()
        except BaseException as error:
            cleanup_failures.append(f"{label}: {type(error).__name__}")

    if primary_error is None:
        try:
            wait_for_exported_slides(
                raw_directory,
                expected_count=expected_count,
                timeout_seconds=timeout_seconds,
            )
        except BaseException as error:
            primary_error = error

    if primary_error is not None:
        if cleanup_failures:
            primary_error.add_note(
                "PowerPoint cleanup failed: " + ", ".join(cleanup_failures)
            )
        raise primary_error
    if cleanup_failures:
        raise RendererCleanupError(cleanup_failures)


def _powerpoint_worker_entry(
    connection: object,
    source: str,
    raw_directory: str,
    expected_count: int,
    timeout_seconds: float,
) -> None:
    try:
        from win32com.client import GetObject

        _powerpoint_export_lifecycle(
            Path(source),
            Path(raw_directory),
            expected_count=expected_count,
            timeout_seconds=timeout_seconds,
            get_object=GetObject,
        )
    except BaseException as error:
        connection.send(
            {
                "kind": "error",
                "primary_type": type(error).__name__,
                "primary_message": str(error),
                "notes": list(getattr(error, "__notes__", ())),
            }
        )
        connection.close()
        raise SystemExit(1)
    else:
        connection.send({"kind": "complete"})
        connection.close()


def _worker_has_started(process: object) -> bool:
    try:
        pid = getattr(process, "pid")
    except (AttributeError, ValueError):
        pid = None
    if pid is not None:
        return True
    try:
        getattr(process, "sentinel")
    except (AttributeError, ValueError):
        return False
    return True


def _stop_worker(process: object) -> None:
    failures: list[str] = []

    try:
        alive = bool(process.is_alive())
    except BaseException as error:
        failures.append(f"worker liveness: {type(error).__name__}")
        alive = True

    if alive:
        try:
            process.terminate()
        except BaseException as error:
            failures.append(f"worker terminate: {type(error).__name__}")
        try:
            process.join(1)
        except BaseException as error:
            failures.append(f"worker join after terminate: {type(error).__name__}")
    else:
        try:
            process.join(0)
        except BaseException as error:
            failures.append(f"worker join: {type(error).__name__}")

    try:
        alive = bool(process.is_alive())
    except BaseException as error:
        failures.append(f"worker liveness after terminate: {type(error).__name__}")
        alive = True

    if alive:
        try:
            process.kill()
        except BaseException as error:
            failures.append(f"worker kill: {type(error).__name__}")
        try:
            process.join(1)
        except BaseException as error:
            failures.append(f"worker join after kill: {type(error).__name__}")

    try:
        if process.is_alive():
            failures.append("worker still alive after kill")
    except BaseException as error:
        failures.append(f"worker final liveness: {type(error).__name__}")

    if failures:
        raise RendererCleanupError(failures)


def _stop_owned_external(process: subprocess.Popen[bytes]) -> None:
    failures: list[str] = []

    try:
        alive = process.poll() is None
    except BaseException as error:
        failures.append(f"owned renderer liveness: {type(error).__name__}")
        alive = True

    if alive:
        try:
            process.terminate()
        except BaseException as error:
            failures.append(f"owned renderer terminate: {type(error).__name__}")
        try:
            process.wait(timeout=2)
        except subprocess.TimeoutExpired:
            pass
        except BaseException as error:
            failures.append(f"owned renderer wait after terminate: {type(error).__name__}")

    try:
        alive = process.poll() is None
    except BaseException as error:
        failures.append(f"owned renderer liveness after terminate: {type(error).__name__}")
        alive = True

    if alive:
        try:
            process.kill()
        except BaseException as error:
            failures.append(f"owned renderer kill: {type(error).__name__}")
        try:
            process.wait(timeout=2)
        except subprocess.TimeoutExpired:
            failures.append("owned renderer still alive after kill")
        except BaseException as error:
            failures.append(f"owned renderer wait after kill: {type(error).__name__}")

    try:
        if process.poll() is None:
            failures.append("owned renderer still alive after cleanup")
    except BaseException as error:
        failures.append(f"owned renderer final liveness: {type(error).__name__}")

    if failures:
        raise RendererCleanupError(failures)


def _run_owned_renderer_process(
    entrypoint: object,
    args: tuple[object, ...],
    *,
    timeout_seconds: float,
    external_launcher: object | None = None,
    ownership_guard: object | None = None,
    completion_check: object | None = None,
) -> ProcessResult:
    if timeout_seconds <= 0:
        raise ValueError("renderer timeout must be positive")
    context = multiprocessing.get_context("spawn")
    receive, send = context.Pipe(duplex=False)
    started_at = datetime.now(timezone.utc)
    started = time.monotonic()
    deadline = started + timeout_seconds
    process = context.Process(target=entrypoint, args=(send, *args))
    process_started = False
    external: subprocess.Popen[bytes] | None = None
    worker_pid: int | None = None
    final_message: dict[str, object] | None = None
    try:
        if external_launcher is not None:
            external = external_launcher()
        process.start()
        process_started = True
        send.close()
        assert process.pid is not None
        worker_pid = process.pid
        while time.monotonic() < deadline:
            remaining = max(0.0, deadline - time.monotonic())
            if receive.poll(min(0.05, remaining)):
                try:
                    message = receive.recv()
                except EOFError:
                    break
                if message.get("kind") in {"complete", "error"}:
                    final_message = message
                    break
            if external is not None:
                if external.poll() is not None:
                    raise RendererLifecycleError(
                        "RendererOwnershipLost",
                        "Tiến trình renderer được tạo trực tiếp đã kết thúc trước worker.",
                    )
                if ownership_guard is not None and not ownership_guard(external.pid):
                    raise RendererLifecycleError(
                        "RendererOwnershipContaminated",
                        "Phát hiện PowerPoint không thuộc lượt render; dừng fail-closed.",
                    )
            if not process.is_alive() and not receive.poll():
                break

        if final_message is None and time.monotonic() >= deadline:
            elapsed = time.monotonic() - started
            result = ProcessResult(
                started=True,
                exit_code=process.exitcode,
                timed_out=True,
                elapsed_seconds=elapsed,
                pid=worker_pid or 0,
                started_at=started_at,
                failure_kind="timeout",
            )
            raise RendererTimeoutError(worker_pid or 0, result)

        process.join(max(0.0, deadline - time.monotonic()))
        if process_started and process.is_alive():
            elapsed = time.monotonic() - started
            result = ProcessResult(
                started=True,
                exit_code=process.exitcode,
                timed_out=True,
                elapsed_seconds=elapsed,
                pid=worker_pid or 0,
                started_at=started_at,
                failure_kind="timeout",
            )
            raise RendererTimeoutError(worker_pid or 0, result)

        elapsed = time.monotonic() - started
        if final_message is None:
            result = ProcessResult(
                started=True,
                exit_code=process.exitcode if process.exitcode not in (None, 0) else 1,
                timed_out=False,
                elapsed_seconds=elapsed,
                pid=worker_pid or 0,
                started_at=started_at,
                failure_kind="renderer",
            )
            raise RendererLifecycleError(
                "RendererWorkerExit",
                "Renderer worker kết thúc mà không trả kết quả.",
                process_result=result,
            )
        if final_message.get("kind") == "error":
            result = ProcessResult(
                started=True,
                exit_code=process.exitcode if process.exitcode not in (None, 0) else 1,
                timed_out=False,
                elapsed_seconds=elapsed,
                pid=worker_pid or 0,
                started_at=started_at,
                failure_kind="renderer",
            )
            raise RendererLifecycleError(
                str(final_message.get("primary_type", "RendererError")),
                str(final_message.get("primary_message", "Renderer thất bại.")),
                notes=tuple(str(note) for note in final_message.get("notes", ())),
                process_result=result,
            )
        if external is not None:
            _stop_owned_external(external)
        if completion_check is not None:
            remaining = max(0.0, deadline - time.monotonic())
            try:
                completion_check(remaining)
            except BaseException as error:
                raise RendererLifecycleError(
                    type(error).__name__,
                    str(error),
                ) from error
        return ProcessResult(
            started=True,
            exit_code=0,
            timed_out=False,
            elapsed_seconds=elapsed,
            pid=worker_pid or 0,
            started_at=started_at,
        )
    finally:
        primary_error = sys.exc_info()[1]
        cleanup_failures: list[str] = []
        worker_started = process_started
        if not worker_started:
            try:
                worker_started = _worker_has_started(process)
            except BaseException as error:
                cleanup_failures.append(f"worker start state: {type(error).__name__}")
        if worker_started:
            try:
                _stop_worker(process)
            except RendererCleanupError as error:
                cleanup_failures.extend(error.failures)
            except BaseException as error:
                cleanup_failures.append(f"worker cleanup: {type(error).__name__}")
        if external is not None:
            try:
                _stop_owned_external(external)
            except RendererCleanupError as error:
                cleanup_failures.extend(error.failures)
            except BaseException as error:
                cleanup_failures.append(f"owned renderer cleanup: {type(error).__name__}")
        try:
            receive.close()
        except BaseException as error:
            cleanup_failures.append(f"worker pipe: {type(error).__name__}")
        try:
            send.close()
        except BaseException as error:
            cleanup_failures.append(f"worker send pipe: {type(error).__name__}")
        if cleanup_failures:
            if primary_error is not None:
                primary_error.add_note(
                    "renderer supervisor cleanup failed: "
                    + ", ".join(cleanup_failures)
                )
            else:
                raise RendererCleanupError(cleanup_failures)


def _run_powerpoint_renderer(
    source: Path,
    raw_directory: Path,
    expected_count: int,
    timeout_seconds: float,
) -> ProcessResult:
    discovery = discover_renderer()
    if not discovery.available or discovery.executable is None:
        raise RendererLifecycleError(
            "RendererUnavailable",
            discovery.unavailable_reason or "Không tìm thấy PowerPoint executable.",
        )
    if _powerpoint_pids():
        raise RendererLifecycleError(
            "RendererOwnershipUnproven",
            "Đã có PowerPoint chạy trước lượt render; không thể chứng minh ownership riêng.",
        )

    def launch() -> subprocess.Popen[bytes]:
        return subprocess.Popen(
            [str(discovery.executable), "/automation", str(source.resolve())],
            stdin=subprocess.DEVNULL,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )

    def ownership_guard(owned_pid: int) -> bool:
        return not (_powerpoint_pids() - {owned_pid})

    return _run_owned_renderer_process(
        _powerpoint_worker_entry,
        (str(source), str(raw_directory), expected_count, timeout_seconds),
        timeout_seconds=timeout_seconds,
        external_launcher=launch,
        ownership_guard=ownership_guard,
        completion_check=lambda remaining: wait_for_exported_slides(
            raw_directory,
            expected_count=expected_count,
            timeout_seconds=remaining,
        ),
    )


def wait_for_exported_slides(
    directory: Path,
    *,
    expected_count: int,
    timeout_seconds: float,
    poll_interval_seconds: float = 0.05,
) -> list[Path]:
    """Đợi PowerPoint hoàn tất ghi đủ ảnh, dựa trên trạng thái file thực tế."""
    deadline = time.monotonic() + timeout_seconds
    previous_sizes: tuple[int, ...] | None = None
    while time.monotonic() < deadline:
        exported = sorted(
            directory.glob("*.PNG"),
            key=lambda path: int(re.search(r"(\d+)$", path.stem).group(1)),
        )
        identities = tuple(path.lstat() for path in exported)
        sizes = tuple(info.st_size for info in identities)
        reparse_marker = getattr(stat, "FILE_ATTRIBUTE_REPARSE_POINT", 0)
        identities_are_safe = all(
            stat.S_ISREG(info.st_mode)
            and info.st_nlink == 1
            and not bool(getattr(info, "st_file_attributes", 0) & reparse_marker)
            for info in identities
        )
        if (
            len(exported) == expected_count
            and all(size > 0 for size in sizes)
            and identities_are_safe
            and sizes == previous_sizes
        ):
            return exported
        previous_sizes = sizes
        time.sleep(poll_interval_seconds)
    raise RuntimeError(
        f"Renderer không hoàn tất {expected_count} ảnh trong {timeout_seconds:g} giây."
    )


def render_staging_path(render_root: Path, token: str) -> Path:
    """Tạo đường staging không có suffix vì PowerPoint sẽ tự bỏ suffix."""
    return render_root / f"staging-{token}"


def _read_bound_file(binding: BoundDirectory, name: str) -> bytes:
    descriptor = binding.open_file(name, os.O_RDONLY)
    payload = bytearray()
    primary_error: BaseException | None = None
    try:
        while True:
            chunk = os.read(descriptor, 1024 * 1024)
            if not chunk:
                break
            payload.extend(chunk)
        binding.verify_file(name, descriptor)
    except BaseException as error:
        primary_error = error
    try:
        os.close(descriptor)
    except BaseException as close_error:
        if primary_error is not None:
            primary_error.add_note(f"bound file close failed: {type(close_error).__name__}")
        else:
            raise
    if primary_error is not None:
        raise primary_error
    return bytes(payload)


def _write_bound_file(binding: BoundDirectory, name: str, payload: bytes) -> None:
    descriptor = binding.open_file(
        name, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600
    )
    primary_error: BaseException | None = None
    try:
        view = memoryview(payload)
        while view:
            written = os.write(descriptor, view)
            if written <= 0:
                raise OSError("short render artifact write")
            view = view[written:]
        os.fsync(descriptor)
        binding.verify_file(name, descriptor)
    except BaseException as error:
        primary_error = error
    try:
        os.close(descriptor)
    except BaseException as close_error:
        if primary_error is not None:
            primary_error.add_note(f"bound file close failed: {type(close_error).__name__}")
        else:
            raise
    if primary_error is not None:
        raise primary_error


def _raw_exports(binding: BoundDirectory, expected_count: int) -> list[str]:
    binding.verify()
    names = [entry.name for entry in os.scandir(binding.path)]
    parsed: list[tuple[int, str]] = []
    for name in names:
        match = re.fullmatch(r"Slide(\d+)\.PNG", name, re.IGNORECASE)
        if match is None:
            raise RenderArtifactError(f"Renderer trả artifact không mong đợi: {name}")
        payload = _read_bound_file(binding, name)
        if not payload:
            raise RenderArtifactError(f"Renderer trả artifact rỗng: {name}")
        parsed.append((int(match.group(1)), name))
    parsed.sort()
    if [index for index, _ in parsed] != list(range(1, expected_count + 1)):
        raise RenderArtifactError(
            f"Renderer trả {len(parsed)} artifact, không khớp {expected_count} slide."
        )
    binding.verify()
    return [name for _, name in parsed]


def _validate_render_artifacts(
    project_root: Path,
    binding: BoundDirectory,
    slides: dict[str, RenderedSlide],
) -> None:
    binding.verify()
    actual_names = sorted(entry.name for entry in os.scandir(binding.path))
    expected_names = sorted(Path(item.path).name for item in slides.values())
    if actual_names != expected_names:
        raise RenderArtifactError("Danh sách artifact sau promote không khớp kết quả render.")
    for slide_id, rendered in slides.items():
        target = safe_path(project_root, rendered.path)
        if target.parent.absolute() != binding.path.absolute():
            raise RenderArtifactError(
                f"Đường artifact sau promote nằm ngoài namespace của slide {slide_id}."
            )
        payload = _read_bound_file(binding, target.name)
        digest = hashlib.sha256(payload).hexdigest()
        if digest != rendered.sha256:
            raise RenderArtifactError(
                f"Hash artifact sau promote không khớp slide {slide_id}."
            )
    binding.verify()


def _record_cleanup_failure(
    primary_error: BaseException, label: str, action: object
) -> None:
    try:
        action()
    except BaseException as cleanup_error:
        primary_error.add_note(f"{label} cleanup failed: {type(cleanup_error).__name__}")


def _remove_stage_when_safe(stage: _BoundStage, timeout_seconds: float = 2.0) -> None:
    deadline = time.monotonic() + timeout_seconds
    while True:
        try:
            stage.remove()
            return
        except UnsafeFileError:
            if time.monotonic() >= deadline:
                raise
            time.sleep(0.05)


def _semantic_issue(slide_id: str, element_id: str | None, detail: str) -> LayoutIssue:
    return LayoutIssue(
        code="SEMANTIC_MISMATCH",
        message_vi=f"PPTX mở lại không khớp nội dung canonical: {detail}",
        slide_id=slide_id,
        element_id=element_id,
        suggested_action_vi="Không promote artifact; kiểm tra adapter native và chạy lại.",
    )


def _verify_semantics(build_input: BuildInput, actual: list[SlideSemantic]) -> list[LayoutIssue]:
    deck = build_input.deck
    if [slide.slide_id for slide in actual] != [slide.slide_id for slide in deck.slides]:
        return [_semantic_issue(deck.slides[0].slide_id, None, "thứ tự hoặc số lượng slide")]
    issues: list[LayoutIssue] = []
    for expected, observed in zip(deck.slides, actual):
        if (observed.title, observed.message, observed.notes, observed.source_refs) != (
            expected.title,
            expected.message,
            expected.notes,
            expected.source_refs,
        ):
            issues.append(_semantic_issue(expected.slide_id, None, "title/message/notes/source_refs"))
            continue
        if list(observed.elements) != [element.element_id for element in expected.elements]:
            issues.append(_semantic_issue(expected.slide_id, None, "thứ tự hoặc số lượng phần tử"))
            continue
        for element in expected.elements:
            value = observed.elements[element.element_id]
            if isinstance(element, TextElement):
                matches = value.kind == "text" and value.text == _format_text(element)
            elif isinstance(element, TableElement):
                expected_rows = [
                    ["" if item is None else str(item) for item in row]
                    for row in element.content.rows
                ]
                matches = (
                    value.kind == "table"
                    and value.columns == element.content.columns
                    and value.rows == expected_rows
                )
            elif isinstance(element, ChartElement):
                matches = (
                    value.kind == "chart"
                    and value.chart_type == element.content.chart_type
                    and value.data
                    == [(datum.label, float(datum.value)) for datum in element.content.data]
                    and value.unit == element.content.unit
                    and value.source_ref == element.content.source_ref
                )
            elif isinstance(element, ProcessElement):
                matches = value.kind == "process" and value.text == _format_process(element)
            elif isinstance(element, TimelineElement):
                matches = value.kind == "timeline" and value.text == _format_timeline(element)
            elif isinstance(element, QuoteElement):
                matches = value.kind == "quote" and value.text == _format_quote(element)
            else:
                matches = False
            if not matches:
                issues.append(
                    _semantic_issue(expected.slide_id, element.element_id, element.kind)
                )
    return issues


class PptxNativeBackend(PresentationBackend):
    def __init__(
        self,
        *,
        renderer: RendererDiscovery | None = None,
        renderer_runner: object | None = None,
    ):
        self.renderer = renderer or discover_renderer()
        self.renderer_runner = renderer_runner or _run_powerpoint_renderer

    def capabilities(self) -> BackendCapabilities:
        return native_capabilities(self.renderer)

    def build(self, build_input: BuildInput, paths: WorkspacePaths) -> BuildResult:
        if not _identity_safe_native_promotion_available():
            raise RuntimeError(
                "Backend native chưa có identity-safe promotion đã kiểm chứng "
                "trên nền tảng này."
            )
        issues = preflight_deck(build_input.deck)
        if issues:
            raise NativeBuildError(issues)

        started_at = datetime.now(timezone.utc)
        started = time.perf_counter()
        build_id = f"native-{build_input.input_hash[:12]}-{uuid4().hex[:8]}"
        builds_root = safe_path(
            paths.project_root, paths.builds_dir.relative_to(paths.project_root)
        )
        builds_root.mkdir(parents=True, exist_ok=True)
        final_dir = safe_path(paths.project_root, f"builds/{build_id}")
        stage = _BoundStage.create(builds_root, f"staging-{build_id}-{uuid4().hex}")
        staged_output = stage.path / "deck.pptx"
        relative_output = f"builds/{build_id}/deck.pptx"

        try:
            presentation = Presentation()
            width, height = canvas_inches(build_input.deck.canvas)
            presentation.slide_width = Inches(width)
            presentation.slide_height = Inches(height)
            presentation.core_properties.title = build_input.deck.title
            presentation.core_properties.subject = build_input.deck.purpose
            presentation.core_properties.author = "Agent Presentation Studio"
            for slide in build_input.deck.slides:
                _render_slide(presentation, slide, width, height)
            serialized = io.BytesIO()
            presentation.save(serialized)
            _write_bound_file(stage.directory, "deck.pptx", serialized.getvalue())

            stage.verify()
            semantic = extract_semantics(
                io.BytesIO(_read_bound_file(stage.directory, "deck.pptx"))
            )
            stage.verify()
            semantic_issues = _verify_semantics(build_input, semantic)
            if semantic_issues:
                raise NativeBuildError(semantic_issues)
            digest = hashlib.sha256(_read_bound_file(stage.directory, "deck.pptx")).hexdigest()
            warning = (
                "Chưa kiểm chứng hình ảnh: renderer đã được phát hiện nhưng chưa chạy."
                if self.renderer.available
                else f"Chưa kiểm chứng hình ảnh: {self.renderer.unavailable_reason}"
            )
            slide_results = [
                SlideBuildResult(
                    slide_id=slide.slide_id,
                    status="passed",
                    artifact_paths=[relative_output],
                    element_editability={
                        element.element_id: "native" for element in slide.elements
                    },
                )
                for slide in build_input.deck.slides
            ]
            result = BuildResult(
                build_id=build_id,
                backend="pptx-native",
                status="passed",
                input_hash=build_input.input_hash,
                outputs=[
                    OutputArtifact(
                        path=relative_output,
                        media_type=PPTX_MEDIA_TYPE,
                        sha256=digest,
                        slide_count=len(semantic),
                        editability="native",
                        notes_included=any(slide.notes is not None for slide in build_input.deck.slides),
                        limitations=[warning],
                    )
                ],
                slide_results=slide_results,
                expected_slide_ids=[slide.slide_id for slide in build_input.deck.slides],
                actual_slide_ids=[slide.slide_id for slide in semantic],
                warnings=[warning],
                process_result=ProcessResult(
                    started=True,
                    exit_code=0,
                    timed_out=False,
                    elapsed_seconds=time.perf_counter() - started,
                    pid=os.getpid(),
                    started_at=started_at,
                ),
            )
            stage.promote(final_dir)
            promoted_digest = hashlib.sha256(
                _read_bound_file(stage.directory, "deck.pptx")
            ).hexdigest()
            if promoted_digest != digest:
                raise RenderArtifactError("Hash PPTX sau promote không khớp staged artifact.")
        except BaseException as primary_error:
            _record_cleanup_failure(primary_error, "build staging removal", stage.remove)
            _record_cleanup_failure(primary_error, "build staging close", stage.close)
            raise
        stage.close()
        return result

    def render(
        self,
        paths: WorkspacePaths,
        build_id: str,
        *,
        timeout_seconds: float | None = None,
    ) -> RenderResult:
        if not self.renderer.available or self.renderer.renderer_id != "powerpoint-com":
            raise RuntimeError(self.renderer.unavailable_reason or "Renderer không khả dụng.")
        if _IDENTIFIER.fullmatch(build_id) is None:
            raise ValueError("build_id không hợp lệ")
        source = safe_path(paths.project_root, f"builds/{build_id}/deck.pptx")
        if not source.is_file():
            raise FileNotFoundError(source)
        if timeout_seconds is None:
            timeout_seconds = load_project(paths.project_root).limits.render_timeout_seconds
        if timeout_seconds <= 0:
            raise ValueError("render timeout must be positive")
        semantic = extract_semantics(source)
        render_root = safe_path(paths.project_root, f"builds/{build_id}/renders")
        render_root.mkdir(parents=False, exist_ok=True)
        render_token = uuid4().hex
        staging = render_staging_path(render_root, render_token)
        final_dir = render_root / f"render-{render_token}"
        stage: _BoundStage | None = None
        raw: BoundDirectory | None = None
        mapped: BoundDirectory | None = None
        raw_closed = False
        mapped_closed = False
        mapped_promoted = False
        stage_removed = False
        stage_closed = False
        try:
            stage = _BoundStage.create(render_root, staging.name)
            raw = stage.directory.child("raw", create=True)
            mapped = stage.directory.child("mapped", create=True, allow_delete=True)
            self.renderer_runner(
                source,
                raw.path,
                len(semantic),
                timeout_seconds,
            )
            stage.verify()
            raw_names = _raw_exports(raw, len(semantic))
            slides: dict[str, RenderedSlide] = {}
            dimensions: RenderDimensions | None = None
            for raw_name, slide in zip(raw_names, semantic):
                payload = _read_bound_file(raw, raw_name)
                destination_name = f"{slide.slide_id}.png"
                _write_bound_file(mapped, destination_name, payload)
                with Image.open(io.BytesIO(payload)) as image:
                    current = RenderDimensions(width=image.width, height=image.height)
                dimensions = dimensions or current
                if current != dimensions:
                    raise RuntimeError("Renderer trả kích thước slide không nhất quán.")
                relative = (final_dir / destination_name).relative_to(
                    paths.project_root
                ).as_posix()
                slides[slide.slide_id] = RenderedSlide(
                    path=relative,
                    sha256=hashlib.sha256(payload).hexdigest(),
                )
            mapped.rename_noreplace(final_dir)
            mapped_promoted = True
            raw.close()
            raw_closed = True
            _remove_stage_when_safe(stage)
            stage_removed = True
            stage.close()
            stage_closed = True
            _validate_render_artifacts(paths.project_root, mapped, slides)
            result = RenderResult(
                build_id=build_id,
                status="passed",
                slides=slides,
                renderer_version=self.renderer.version,
                dimensions=dimensions,
            )
            mapped.close()
            mapped_closed = True
            return result
        except BaseException as primary_error:
            if raw is not None and not raw_closed:
                _record_cleanup_failure(primary_error, "raw namespace close", raw.close)
            if mapped is not None and mapped_promoted:
                _record_cleanup_failure(
                    primary_error, "promoted render removal", mapped.remove_tree
                )
            if mapped is not None and not mapped_closed:
                _record_cleanup_failure(primary_error, "mapped namespace close", mapped.close)
            if stage is not None and not stage_removed:
                _record_cleanup_failure(
                    primary_error,
                    "render staging removal",
                    lambda: _remove_stage_when_safe(stage),
                )
            if stage is not None and not stage_closed:
                _record_cleanup_failure(primary_error, "render staging close", stage.close)
            raise


def _cli_errors(issues: tuple[LayoutIssue, ...]) -> list[CLIError]:
    return [
        CLIError(
            code=issue.code,
            message_vi=issue.message_vi,
            field=(f"elements.{issue.element_id}" if issue.element_id else "layout_ref"),
            slide_id=issue.slide_id,
        )
        for issue in issues
    ]


def build_workspace(paths: WorkspacePaths) -> CLIResult:
    project = load_project(paths.project_root)
    deck = load_deck(paths.project_root, project.deck_path)
    profile_lock = resolve_profile(project.profile_ref, paths) if project.profile_ref else None
    payload = {
        "project": project.model_dump(mode="json"),
        "deck": deck.model_dump(mode="json"),
        "profile_lock": profile_lock.model_dump(mode="json") if profile_lock else None,
    }
    build_input = BuildInput(
        project=project,
        deck=deck,
        profile_lock=profile_lock,
        input_hash=hash_inputs(payload),
    )
    backend = PptxNativeBackend()
    try:
        result = backend.build(build_input, paths)
    except NativeBuildError as exc:
        return CLIResult(
            command="build",
            status="failed",
            exit_code=2,
            errors=_cli_errors(exc.issues),
        )
    capabilities = backend.capabilities()
    capabilities.verification_by_feature = {
        **capabilities.verification_by_feature,
        "semantic-reopen": FeatureVerification(
            status="passed",
            evidence_ref=result.outputs[0].path,
            environment=f"python-pptx {pptx.__version__}",
        ),
    }
    return CLIResult(
        command="build",
        status="unverified",
        exit_code=0,
        data={
            "build": result.model_dump(mode="json"),
            "capabilities": capabilities.model_dump(mode="json"),
            "visual_status": "unverified",
        },
        warnings=result.warnings,
    )


def render_workspace(paths: WorkspacePaths, build_id: str) -> CLIResult:
    backend = PptxNativeBackend()
    if not backend.renderer.available:
        return CLIResult(
            command="render",
            status="failed",
            exit_code=3,
            errors=[
                CLIError(
                    code="CAPABILITY_UNAVAILABLE",
                    message_vi=backend.renderer.unavailable_reason or "Renderer không khả dụng.",
                )
            ],
        )
    project = load_project(paths.project_root)
    try:
        result = backend.render(
            paths,
            build_id,
            timeout_seconds=project.limits.render_timeout_seconds,
        )
    except RendererTimeoutError as exc:
        return CLIResult(
            command="render",
            status="failed",
            exit_code=5,
            data={"process_result": exc.process_result.model_dump(mode="json")},
            errors=[
                CLIError(
                    code="RENDER_TIMEOUT",
                    message_vi="Renderer vượt quá giới hạn thời gian và đã được hủy an toàn.",
                )
            ],
        )
    except (RendererLifecycleError, RendererCleanupError, RenderArtifactError) as exc:
        notes = list(getattr(exc, "__notes__", ()))
        if isinstance(exc, RendererCleanupError):
            notes.extend(exc.failures)
        process_result = getattr(exc, "process_result", None)
        return CLIResult(
            command="render",
            status="failed",
            exit_code=5,
            data={
                "cleanup_evidence": notes,
                "process_result": (
                    process_result.model_dump(mode="json")
                    if process_result is not None
                    else None
                ),
            },
            errors=[
                CLIError(
                    code="RENDERER_FAILURE",
                    message_vi="Renderer thất bại; artifact chưa được công bố là đạt.",
                )
            ],
        )
    return CLIResult(
        command="render",
        status="passed",
        exit_code=0,
        data={"render": result.model_dump(mode="json")},
    )


__all__ = [
    "NativeBuildError",
    "PptxNativeBackend",
    "RenderArtifactError",
    "RendererCleanupError",
    "RendererDiscovery",
    "RendererLifecycleError",
    "RendererTimeoutError",
    "SlideSemantic",
    "discover_renderer",
    "extract_semantics",
    "native_capabilities",
    "render_staging_path",
    "wait_for_exported_slides",
]
