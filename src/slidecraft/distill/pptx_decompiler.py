"""PPTX Decompiler & Property Inspector for Reverse Distillation."""
import os
from typing import Dict, Any, List
from pptx import Presentation
from slidecraft.ir.theme_schema import ThemeSchema, CanvasConfig, PaletteConfig, TypographyConfig


class PPTXDecompiler:
    """Inspects an existing .pptx file and extracts theme tokens, dimensions, and typography."""

    @classmethod
    def inspect(cls, pptx_path: str, profile_name: str) -> ThemeSchema:
        if not os.path.exists(pptx_path):
            raise FileNotFoundError(f"File not found: {pptx_path}")

        prs = Presentation(pptx_path)
        
        # 1. Canvas Dimensions
        width_inch = prs.slide_width.inches
        height_inch = prs.slide_height.inches
        
        canvas = CanvasConfig(
            width_inch=round(width_inch, 3),
            height_inch=round(height_inch, 3)
        )
        
        # 2. Extract Fonts and Colors from Shapes
        discovered_fonts = set()
        discovered_colors = set()
        
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            if run.font.name:
                                discovered_fonts.add(run.font.name)
                            try:
                                if run.font.color and hasattr(run.font.color, "rgb") and run.font.color.rgb is not None:
                                    discovered_colors.add(f"#{str(run.font.color.rgb)}")
                            except Exception:
                                pass

        primary_font = list(discovered_fonts)[0] if discovered_fonts else "Segoe UI"
        
        typography = TypographyConfig(
            title_font=primary_font,
            body_font=primary_font
        )
        
        palette = PaletteConfig()
        if discovered_colors:
            colors_list = list(discovered_colors)
            palette.primary_text = colors_list[0]
            if len(colors_list) > 1:
                palette.accent_primary = colors_list[1]
                
        return ThemeSchema(
            profile_name=profile_name,
            display_name=f"Distilled: {profile_name}",
            description=f"Auto-extracted from {os.path.basename(pptx_path)}",
            canvas=canvas,
            palette=palette,
            typography=typography
        )
