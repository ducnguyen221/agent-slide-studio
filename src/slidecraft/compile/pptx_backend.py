"""Native PPTX Compilation Backend for Slidecraft."""
import os
from typing import List, Optional
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

from slidecraft.ir.deck import DeckSchema, SlideSchema, SlideLayoutArchetype
from slidecraft.ir.theme_schema import ThemeSchema
from slidecraft.ir.components import CardNode, MetricNode, ProcessStepNode, TimelineNode
from slidecraft.layout.constraint_solver import ConstraintSolver, BoundingBox
from slidecraft.compile.ooxml_enhancer import hex_to_rgb, style_text_frame


class PPTXBackend:
    """Compiles Canonical Slidecraft IR into a 100% Native Editable PPTX file."""

    def __init__(self, theme: ThemeSchema):
        self.theme = theme
        self.solver = ConstraintSolver(theme.canvas)

    def create_presentation(self) -> Presentation:
        prs = Presentation()
        prs.slide_width = Inches(self.theme.canvas.width_inch)
        prs.slide_height = Inches(self.theme.canvas.height_inch)
        return prs

    def compile_deck(self, deck: DeckSchema, output_path: str) -> str:
        prs = self.create_presentation()
        blank_slide_layout = prs.slide_layouts[6]  # Completely blank layout

        for slide_data in deck.slides:
            slide = prs.slides.add_slide(blank_slide_layout)
            self._apply_background(slide)
            self._render_header(slide, slide_data)
            self._render_body_by_archetype(slide, slide_data)

        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        prs.save(output_path)
        return output_path

    def _apply_background(self, slide):
        """Set slide solid background color."""
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = hex_to_rgb(self.theme.palette.background)

    def _render_header(self, slide, slide_data: SlideSchema):
        """Render category badge and main slide title."""
        header_box = self.solver.get_header_box()
        
        # 1. Category Badge (if present)
        current_top = header_box.top_inch
        if slide_data.category_badge:
            badge_width = 2.4
            badge_height = 0.32
            badge = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(header_box.left_inch),
                Inches(current_top),
                Inches(badge_width),
                Inches(badge_height)
            )
            badge.fill.solid()
            badge.fill.fore_color.rgb = hex_to_rgb(self.theme.palette.muted_badge)
            badge.line.color.rgb = hex_to_rgb(self.theme.palette.surface_card_border)
            badge.line.width = Pt(0.75)
            
            tf = badge.text_frame
            style_text_frame(tf, margin_inch=0.04)
            p = tf.paragraphs[0]
            p.text = slide_data.category_badge.upper()
            p.font.name = self.theme.typography.title_font
            p.font.size = Pt(9.5)
            p.font.bold = True
            p.font.color.rgb = hex_to_rgb(self.theme.palette.accent_primary)
            p.alignment = PP_ALIGN.CENTER
            current_top += badge_height + 0.10

        # 2. Main Title Text Box
        title_box = slide.shapes.add_textbox(
            Inches(header_box.left_inch),
            Inches(current_top),
            Inches(header_box.width_inch),
            Inches(0.65)
        )
        tf = title_box.text_frame
        style_text_frame(tf, margin_inch=0.0)
        p = tf.paragraphs[0]
        p.text = slide_data.title
        p.font.name = self.theme.typography.title_font
        p.font.size = Pt(self.theme.typography.title_size_pt)
        p.font.bold = True
        p.font.color.rgb = hex_to_rgb(self.theme.palette.primary_text)
        
        # Subtitle (if present)
        if slide_data.subtitle:
            p_sub = tf.add_paragraph()
            p_sub.text = slide_data.subtitle
            p_sub.font.name = self.theme.typography.body_font
            p_sub.font.size = Pt(self.theme.typography.subtitle_size_pt)
            p_sub.font.color.rgb = hex_to_rgb(self.theme.palette.secondary_text)

    def _render_body_by_archetype(self, slide, slide_data: SlideSchema):
        """Dispatch slide body rendering to specific archetype handlers."""
        archetype = slide_data.archetype
        
        if archetype == SlideLayoutArchetype.CARDS_GRID:
            self._render_cards_grid(slide, slide_data)
        elif archetype == SlideLayoutArchetype.KPI_GRID:
            self._render_kpi_grid(slide, slide_data)
        elif archetype == SlideLayoutArchetype.PROCESS_FLOW:
            self._render_process_flow(slide, slide_data)
        elif archetype == SlideLayoutArchetype.HORIZONTAL_TIMELINE:
            self._render_timeline(slide, slide_data)
        elif archetype == SlideLayoutArchetype.SPLIT_CONTENT_VISUAL:
            self._render_split_visual(slide, slide_data)
        elif archetype == SlideLayoutArchetype.COMPARISON_2COL:
            self._render_comparison_2col(slide, slide_data)
        else:
            self._render_cards_grid(slide, slide_data)

    def _render_cards_grid(self, slide, slide_data: SlideSchema):
        cards = slide_data.cards or []
        num_cards = max(1, len(cards))
        boxes = self.solver.calculate_columns(num_cards)
        
        for i, card_data in enumerate(cards):
            if i < len(boxes):
                self._draw_card(slide, boxes[i], card_data)

    def _draw_card(self, slide, box: BoundingBox, card: CardNode):
        """Draws a native editable rounded rectangle card with title and bullet points."""
        card_shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(box.left_inch),
            Inches(box.top_inch),
            Inches(box.width_inch),
            Inches(box.height_inch)
        )
        card_shape.fill.solid()
        card_shape.fill.fore_color.rgb = hex_to_rgb(self.theme.palette.surface_card)
        
        border_hex = card.accent_color or self.theme.palette.surface_card_border
        card_shape.line.color.rgb = hex_to_rgb(border_hex)
        card_shape.line.width = Pt(1.25)
        
        tf = card_shape.text_frame
        style_text_frame(tf, margin_inch=0.25)
        
        # Card Headline
        p_title = tf.paragraphs[0]
        p_title.text = card.title
        p_title.font.name = self.theme.typography.title_font
        p_title.font.size = Pt(17.0)
        p_title.font.bold = True
        p_title.font.color.rgb = hex_to_rgb(self.theme.palette.primary_text)
        p_title.space_after = Pt(10.0)
        
        # Subtitle (if any)
        if card.subtitle:
            p_sub = tf.add_paragraph()
            p_sub.text = card.subtitle
            p_sub.font.name = self.theme.typography.body_font
            p_sub.font.size = Pt(12.0)
            p_sub.font.italic = True
            p_sub.font.color.rgb = hex_to_rgb(self.theme.palette.secondary_text)
            p_sub.space_after = Pt(8.0)
            
        # Bullet Points
        for pt in card.points:
            p_pt = tf.add_paragraph()
            p_pt.text = f"•  {pt}"
            p_pt.font.name = self.theme.typography.body_font
            p_pt.font.size = Pt(self.theme.typography.body_size_pt)
            p_pt.font.color.rgb = hex_to_rgb(self.theme.palette.secondary_text)
            p_pt.space_after = Pt(6.0)

    def _render_kpi_grid(self, slide, slide_data: SlideSchema):
        metrics = slide_data.metrics or []
        num_metrics = max(1, len(metrics))
        boxes = self.solver.calculate_columns(num_metrics)
        
        for i, metric in enumerate(metrics):
            if i < len(boxes):
                box = boxes[i]
                shape = slide.shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE,
                    Inches(box.left_inch),
                    Inches(box.top_inch),
                    Inches(box.width_inch),
                    Inches(box.height_inch)
                )
                shape.fill.solid()
                shape.fill.fore_color.rgb = hex_to_rgb(self.theme.palette.surface_card)
                shape.line.color.rgb = hex_to_rgb(self.theme.palette.accent_primary)
                shape.line.width = Pt(1.5)
                
                tf = shape.text_frame
                style_text_frame(tf, margin_inch=0.25)
                
                # Big KPI Number
                p_val = tf.paragraphs[0]
                p_val.text = metric.value
                p_val.font.name = self.theme.typography.title_font
                p_val.font.size = Pt(self.theme.typography.kpi_size_pt)
                p_val.font.bold = True
                p_val.font.color.rgb = hex_to_rgb(self.theme.palette.accent_primary)
                p_val.space_after = Pt(6.0)
                
                # Metric Label
                p_lbl = tf.add_paragraph()
                p_lbl.text = metric.label
                p_lbl.font.name = self.theme.typography.title_font
                p_lbl.font.size = Pt(15.0)
                p_lbl.font.bold = True
                p_lbl.font.color.rgb = hex_to_rgb(self.theme.palette.primary_text)
                p_lbl.space_after = Pt(4.0)
                
                if metric.subtitle:
                    p_sub = tf.add_paragraph()
                    p_sub.text = metric.subtitle
                    p_sub.font.name = self.theme.typography.body_font
                    p_sub.font.size = Pt(self.theme.typography.caption_size_pt)
                    p_sub.font.color.rgb = hex_to_rgb(self.theme.palette.secondary_text)

    def _render_process_flow(self, slide, slide_data: SlideSchema):
        steps = slide_data.process_steps or []
        num_steps = max(1, len(steps))
        boxes = self.solver.calculate_columns(num_steps)
        
        for i, step in enumerate(steps):
            if i < len(boxes):
                box = boxes[i]
                step_shape = slide.shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE,
                    Inches(box.left_inch),
                    Inches(box.top_inch),
                    Inches(box.width_inch),
                    Inches(box.height_inch)
                )
                step_shape.fill.solid()
                step_shape.fill.fore_color.rgb = hex_to_rgb(self.theme.palette.surface_card)
                step_shape.line.color.rgb = hex_to_rgb(self.theme.palette.accent_primary if step.is_active else self.theme.palette.surface_card_border)
                step_shape.line.width = Pt(1.5 if step.is_active else 1.0)
                
                tf = step_shape.text_frame
                style_text_frame(tf, margin_inch=0.20)
                
                # Step Indicator "STEP 01"
                p_num = tf.paragraphs[0]
                p_num.text = f"STEP {step.step_number:02d}"
                p_num.font.name = self.theme.typography.title_font
                p_num.font.size = Pt(11.0)
                p_num.font.bold = True
                p_num.font.color.rgb = hex_to_rgb(self.theme.palette.accent_primary)
                p_num.space_after = Pt(6.0)
                
                # Step Title
                p_title = tf.add_paragraph()
                p_title.text = step.title
                p_title.font.name = self.theme.typography.title_font
                p_title.font.size = Pt(15.0)
                p_title.font.bold = True
                p_title.font.color.rgb = hex_to_rgb(self.theme.palette.primary_text)
                p_title.space_after = Pt(6.0)
                
                # Step Description
                p_desc = tf.add_paragraph()
                p_desc.text = step.description
                p_desc.font.name = self.theme.typography.body_font
                p_desc.font.size = Pt(self.theme.typography.body_size_pt)
                p_desc.font.color.rgb = hex_to_rgb(self.theme.palette.secondary_text)

    def _render_timeline(self, slide, slide_data: SlideSchema):
        nodes = slide_data.timeline_nodes or []
        num_nodes = max(1, len(nodes))
        boxes = self.solver.calculate_columns(num_nodes)
        
        # Horizontal connector line across timeline
        if boxes:
            line_top = boxes[0].top_inch + 0.40
            line_left = boxes[0].left_inch + 0.2
            line_width = (boxes[-1].right_inch - line_left) - 0.2
            line = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(line_left),
                Inches(line_top),
                Inches(line_width),
                Pt(3.0)
            )
            line.fill.solid()
            line.fill.fore_color.rgb = hex_to_rgb(self.theme.palette.accent_secondary)
            line.line.fill.background()

        for i, node in enumerate(nodes):
            if i < len(boxes):
                box = boxes[i]
                node_shape = slide.shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE,
                    Inches(box.left_inch),
                    Inches(box.top_inch + 0.65),
                    Inches(box.width_inch),
                    Inches(box.height_inch - 0.65)
                )
                node_shape.fill.solid()
                node_shape.fill.fore_color.rgb = hex_to_rgb(self.theme.palette.surface_card)
                node_shape.line.color.rgb = hex_to_rgb(self.theme.palette.surface_card_border)
                
                tf = node_shape.text_frame
                style_text_frame(tf, margin_inch=0.18)
                
                # Time Label Badge
                p_time = tf.paragraphs[0]
                p_time.text = node.time_label.upper()
                p_time.font.name = self.theme.typography.title_font
                p_time.font.size = Pt(11.0)
                p_time.font.bold = True
                p_time.font.color.rgb = hex_to_rgb(self.theme.palette.accent_secondary)
                p_time.space_after = Pt(4.0)
                
                p_title = tf.add_paragraph()
                p_title.text = node.title
                p_title.font.name = self.theme.typography.title_font
                p_title.font.size = Pt(14.0)
                p_title.font.bold = True
                p_title.font.color.rgb = hex_to_rgb(self.theme.palette.primary_text)
                p_title.space_after = Pt(4.0)
                
                p_desc = tf.add_paragraph()
                p_desc.text = node.description
                p_desc.font.name = self.theme.typography.body_font
                p_desc.font.size = Pt(self.theme.typography.body_size_pt)
                p_desc.font.color.rgb = hex_to_rgb(self.theme.palette.secondary_text)

    def _render_split_visual(self, slide, slide_data: SlideSchema):
        content_box, visual_box = self.solver.calculate_split_content_visual(content_ratio=0.55)
        
        # Left Content Side (Cards or Bullet container)
        if slide_data.cards and len(slide_data.cards) > 0:
            self._draw_card(slide, content_box, slide_data.cards[0])
            
        # Right Visual Asset Placeholder / Image
        asset = slide_data.visual_asset
        if asset and asset.file_path and os.path.exists(asset.file_path):
            slide.shapes.add_picture(
                asset.file_path,
                Inches(visual_box.left_inch),
                Inches(visual_box.top_inch),
                Inches(visual_box.width_inch),
                Inches(visual_box.height_inch)
            )
        else:
            # Render a placeholder visual card
            ph = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(visual_box.left_inch),
                Inches(visual_box.top_inch),
                Inches(visual_box.width_inch),
                Inches(visual_box.height_inch)
            )
            ph.fill.solid()
            ph.fill.fore_color.rgb = hex_to_rgb(self.theme.palette.muted_badge)
            ph.line.color.rgb = hex_to_rgb(self.theme.palette.accent_primary)
            ph.line.width = Pt(1.5)
            
            tf = ph.text_frame
            style_text_frame(tf, margin_inch=0.3)
            p = tf.paragraphs[0]
            p.text = f"[VISUAL ASSET: {asset.prompt_subject if asset else 'Conceptual Illustration'}]"
            p.font.name = self.theme.typography.code_font
            p.font.size = Pt(13.0)
            p.font.color.rgb = hex_to_rgb(self.theme.palette.accent_primary)
            p.alignment = PP_ALIGN.CENTER

    def _render_comparison_2col(self, slide, slide_data: SlideSchema):
        cards = slide_data.cards or []
        boxes = self.solver.calculate_columns(2)
        
        for i, card_data in enumerate(cards[:2]):
            self._draw_card(slide, boxes[i], card_data)
